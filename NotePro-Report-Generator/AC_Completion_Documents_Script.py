import pandas as pd
from docxtpl import DocxTemplate
from docx.shared import Inches
import logging
import requests
from io import BytesIO
import os
import sys
from datetime import datetime
from PIL import Image
import time
import argparse
import shutil
import subprocess

# -----------------------------------------------------------
# 1) Parse command-line arguments (including --templates_dir)
# -----------------------------------------------------------
parser = argparse.ArgumentParser(description='Generate Anger Control Completion Documents')
parser.add_argument('--csv_file', required=True, help='Path to the CSV file')
parser.add_argument('--start_date', required=True, help='Start date in YYYY-MM-DD format')
parser.add_argument('--end_date', required=True, help='End date in YYYY-MM-DD format')
parser.add_argument('--clinic_folder', required=True, help='Clinic folder identifier')
parser.add_argument('--templates_dir', required=True, help='Path to the clinic-specific templates folder')
parser.add_argument('--output_dir', required=False, help='Optional path to the output directory')
args = parser.parse_args()

csv_file_path = args.csv_file
start_date_str = args.start_date
end_date_str = args.end_date
clinic_folder = args.clinic_folder
templates_dir = args.templates_dir  # <-- new

# --- env overrides + templates root ---
program_code  = os.environ.get('NP_PROGRAM_CODE', 'AC')
templates_dir = os.environ.get('NP_TEMPLATES_DIR', templates_dir)
clinic_folder = os.environ.get('NP_CLINIC_FOLDER', clinic_folder)

if os.path.basename(templates_dir) == clinic_folder:
    _TEMPLATES_ROOT = os.path.dirname(templates_dir)
else:
    _TEMPLATES_ROOT = templates_dir

output_dir = args.output_dir

# Setup logging
log_dir = '/home/notesao/NotePro-Report-Generator/logs'
os.makedirs(log_dir, exist_ok=True)
logging.basicConfig(
    filename=os.path.join(log_dir, f'AC_Completion_Documents_{datetime.now().strftime("%Y%m%d")}.log'),
    level=logging.DEBUG,
    format='%(asctime)s [%(levelname)s] %(message)s'
)
logging.info("AC Completion Documents Script started.")

# -----------------------------------------------------------
# 2) Convert start/end dates
# -----------------------------------------------------------
try:
    start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
    end_date = datetime.strptime(end_date_str, '%Y-%m-%d')
except ValueError as e:
    logging.error(f"Invalid date format for start_date or end_date: {e}")
    sys.exit(1)

base_dir = os.path.dirname(os.path.abspath(__file__))
today_str = datetime.now().strftime('%m.%d.%y')

# -----------------------------------------------------------
# 3) Determine output paths
# -----------------------------------------------------------
if output_dir:
    generated_documents_dir = output_dir
else:
    generated_documents_dir = os.path.join(base_dir, 'GeneratedDocuments', clinic_folder, today_str)
os.makedirs(generated_documents_dir, exist_ok=True)

public_generated_documents_dir = f"/home/notesao/{clinic_folder}/public_html/GeneratedDocuments"
os.makedirs(public_generated_documents_dir, exist_ok=True)

central_location_path = os.path.join(public_generated_documents_dir, today_str)
os.makedirs(central_location_path, exist_ok=True)
logging.info(f"Central location directory created: {central_location_path}")

errors_txt = os.path.join(generated_documents_dir, 'errors.txt')
def record_error(msg: str):
    try:
        with open(errors_txt, 'a', encoding='utf-8') as f:
            f.write(str(msg).rstrip() + '\n')
    except Exception:
        pass
    logging.error(msg)

def resolve_template(templates_root, clinic, program, filename):
    for p in (
        os.path.join(templates_root, clinic, filename),
        os.path.join(templates_root, 'default', program, filename),
        os.path.join(templates_root, 'default', 'BIPP', filename),
    ):
        if os.path.isfile(p):
            return p
    return None


# -----------------------------------------------------------
# 4) Template paths via --templates_dir
# -----------------------------------------------------------
completion_letter_template_path      = resolve_template(_TEMPLATES_ROOT, clinic_folder, program_code, 'Template.AC Completion Letter.docx')
completion_certificate_template_path = resolve_template(_TEMPLATES_ROOT, clinic_folder, program_code, 'Template.AC Completion Certificate.docx')
completion_report_template_path      = resolve_template(_TEMPLATES_ROOT, clinic_folder, program_code, 'Template.AC Completion Progress Report.docx')

for _name, _path in [
    ('Template.AC Completion Letter.docx', completion_letter_template_path),
    ('Template.AC Completion Certificate.docx', completion_certificate_template_path),
    ('Template.AC Completion Progress Report.docx', completion_report_template_path),
]:
    if not _path:
        record_error(f"[FATAL] Missing template {_name} under clinic/default fallbacks in {_TEMPLATES_ROOT}.")


# ... rest of your script remains the same ...
def sanitize_filename(filename):
    invalid_chars = '<>:"/\\|?*'
    for char in invalid_chars:
        filename = filename.replace(char, '')
    return filename.strip()

def docx_to_pdf(docx_path, pdf_dir, retries=3):
    """Convert DOCX to PDF using LibreOffice with retry logic."""
    attempt = 0
    while attempt < retries:
        try:
            output_file = os.path.join(pdf_dir, os.path.basename(docx_path).replace('.docx', '.pdf'))
            subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', pdf_dir, docx_path],
                           check=True)
            if os.path.exists(output_file):
                logging.info(f"Converted {docx_path} to PDF successfully on attempt {attempt + 1}.")
                return output_file
            else:
                raise FileNotFoundError(f"PDF not found at expected path: {output_file}")
        except subprocess.CalledProcessError as e:
            logging.error(f"Error converting {docx_path} to PDF on attempt {attempt + 1}: {e}")
            attempt += 1
            time.sleep(2)
    logging.error(f"Failed to convert {docx_path} to PDF after {retries} attempts.")
    return None

def parse_multiple_formats(value, formats=['%m/%d/%Y', '%Y-%m-%d', '%m/%d/%y']):
    """Try several date formats until one succeeds, or return None."""
    if not value or not isinstance(value, str):
        return None
    value = value.strip()
    for fmt in formats:
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None

def fill_placeholders(row):
    """Fill placeholders: replace 'nan', convert date fields, handle attendance, etc."""
    # 0) Replace any 'nan' strings with empty ''
    for col, val in row.items():
        if isinstance(val, str) and val.strip().lower() == 'nan':
            row[col] = ''

    # 1) Convert any field ending in '_date' to mm/dd/yyyy
    for placeholder, value in row.items():
        if placeholder.endswith('_date') and value.strip():
            date_obj = parse_multiple_formats(value)
            if date_obj:
                row[placeholder] = date_obj.strftime('%m/%d/%Y')

    # 2) Specifically handle 'dob' → mm/dd/yyyy
    if 'dob' in row:
        raw_dob = row['dob'].strip()
        if raw_dob:
            date_obj = parse_multiple_formats(raw_dob)
            if date_obj:
                row['dob'] = date_obj.strftime('%m/%d/%Y')
            else:
                row['dob'] = ''
        else:
            row['dob'] = ''

    # 3) ALWAYS set these attendance placeholders for COMPLETION:
    #    R1..R2, C1..C2, S1..S2, O1..O2 → blank
    #    R3, C3, S3, O3 → checkmark
    row['R1'] = ''
    row['R2'] = ''
    row['R3'] = '\u2714'

    row['C1'] = ''
    row['C2'] = ''
    row['C3'] = '\u2714'

    row['S1'] = ''
    row['S2'] = ''
    row['S3'] = '\u2714'

    row['O1'] = ''
    row['O2'] = ''
    row['O3'] = '\u2714'

    # 3) Set ALL D1..D30 to checkmark for COMPLETION
    for i in range(1, 31):
        row[f'D{i}'] = '\u2714'

    # 4) Format P1..P27 as short dates (mm/dd/yy)
    for i in range(1, 28):
        p_field = f'P{i}'
        if p_field in row:
            val = row[p_field].strip()
            if val:
                date_obj = parse_multiple_formats(val)
                if date_obj:
                    row[p_field] = date_obj.strftime('%m/%d/%y')
                else:
                    row[p_field] = ''
            else:
                row[p_field] = ''

    # 5) Format A1..A18 as short dates (mm/dd/yy)
    for i in range(1, 19):
        a_field = f'A{i}'
        if a_field in row:
            val = row[a_field].strip()
            if val:
                date_obj = parse_multiple_formats(val)
                if date_obj:
                    row[a_field] = date_obj.strftime('%m/%d/%y')
                else:
                    row[a_field] = ''
            else:
                row[a_field] = ''

    # 6) Replace '&' with 'and' in client_note
    if 'client_note' in row and isinstance(row['client_note'], str):
        row['client_note'] = row['client_note'].replace('&', 'and')

    # 7) Fee problems
    if 'balance' in row:
        try:
            balance = float(row['balance'])
            row['feeproblems'] = "No" if balance == 0 else "Yes"
        except ValueError:
            row['feeproblems'] = "Yes"

    # 8) If last_attended is empty, fallback to orientation_date
    if 'orientation_date' in row and 'last_attended' in row:
        if not row['last_attended'].strip():
            row['last_attended'] = row['orientation_date']

    return row

def fill_gender_placeholders(context):
    """Fill additional placeholders for gender pronouns."""
    gender = context.get('gender', '').lower()
    if gender == 'male':
        context['gender1'] = 'his'
        context['gender2'] = 'he'
        context['gender3'] = "Men's"
        context['gender4'] = 'him'
        context['gender5'] = 'He'
        context['gender6'] = 'himself'
    elif gender == 'female':
        context['gender1'] = 'her'
        context['gender2'] = 'she'
        context['gender3'] = "Women's"
        context['gender4'] = 'her'
        context['gender5'] = 'She'
        context['gender6'] = 'herself'
    else:
        logging.error(f"Unrecognized gender: {gender}")
        context.update({g: '' for g in ['gender1','gender2','gender3','gender4','gender5','gender6']})

def render_and_save(template_path, doc_type, context):
    """Render a template, convert to PDF, remove the .docx, return PDF path."""
    if not template_path:
        return None  # skip this document type but continue the row

    try:
        doc = DocxTemplate(template_path)
        doc.render(context)

        first_name = sanitize_filename(context.get('first_name', ''))
        last_name = sanitize_filename(context.get('last_name', ''))
        docx_filename = f"{last_name} {first_name} {doc_type}.docx"
        docx_path = os.path.join(generated_documents_dir, docx_filename)

        doc.save(docx_path)
        pdf_path = docx_to_pdf(docx_path, generated_documents_dir)
        if pdf_path:
            os.remove(docx_path)
            return pdf_path
        else:
            logging.error(f"PDF conversion failed for {docx_path}")
            return None
    except Exception as e:
        logging.error(f"Error rendering {doc_type}: {e}")
        return None

# -----------------------------------------------------------
# 5) Load CSV, filter by exit_date
# -----------------------------------------------------------
try:
    df = pd.read_csv(csv_file_path)
    logging.info(f"Loaded {len(df)} records from CSV.")
except Exception as e:
    logging.error(f"Failed to load CSV: {e}")
    sys.exit(1)

df['exit_date'] = pd.to_datetime(df['exit_date'], errors='coerce')
df = df.dropna(subset=['exit_date'])

# Filter for Anger Control + "Completion of Program" + date range
df = df[
    (df['program_name'] == 'Anger Control') &
    (df['exit_reason'] == 'Completion of Program') &
    (df['exit_date'] >= start_date) &
    (df['exit_date'] <= end_date)
]

if df.empty:
    logging.warning("No records found after filtering. No documents will be generated.")
    sys.exit(0)

# Fill placeholders across the DataFrame
df = df.astype(str)
df = df.apply(fill_placeholders, axis=1)
df.fillna('', inplace=True)

def generate_documents_for_row(row):
    """Generate letter, cert, progress report for a single row, then move them to the manager folder."""
    context = row.to_dict()

    # Fill gender placeholders
    fill_gender_placeholders(context)

    # Add today's date as 'report_date'
    now = datetime.now()
    context['report_date'] = f"{now.month}/{now.day}/{now.year}"

    # Render each doc
    letter_pdf = render_and_save(completion_letter_template_path, "AC Comp Letter", context)
    cert_pdf = render_and_save(completion_certificate_template_path, "AC Comp Cert", context)
    progress_pdf = render_and_save(completion_report_template_path, "AC Comp PReport", context)

    case_manager_office = str(context.get('case_manager_office', '')).strip()
    case_manager_first  = str(context.get('case_manager_first_name', '')).strip()
    case_manager_last   = str(context.get('case_manager_last_name', '')).strip()

    # If office is missing or "nan", place them directly in central_location_path
    if not case_manager_office or case_manager_office.lower() == 'nan':
        office_dir = central_location_path
    else:
        office_dir = os.path.join(central_location_path, sanitize_filename(case_manager_office))

    # If first name is missing/'nan', only use last name
    if not case_manager_first or case_manager_first.lower() == 'nan':
        manager_subfolder_name = sanitize_filename(case_manager_last)
    else:
        manager_subfolder_name = f"{sanitize_filename(case_manager_first)} {sanitize_filename(case_manager_last)}"

    manager_subfolder_name = manager_subfolder_name.strip()
    manager_dir = os.path.join(office_dir, manager_subfolder_name)
    os.makedirs(manager_dir, exist_ok=True)

    for pdf_file in [letter_pdf, cert_pdf, progress_pdf]:
        if pdf_file and os.path.exists(pdf_file):
            final_pdf_path = os.path.join(manager_dir, os.path.basename(pdf_file))
            shutil.move(pdf_file, final_pdf_path)
            logging.info(f"Moved {pdf_file} to {final_pdf_path}")
        else:
            logging.error("PDF conversion failed or file doesn't exist; document not moved.")

# Main loop
try:
    for _, row in df.iterrows():
        generate_documents_for_row(row)
except Exception as e:
    logging.error(f"Error during document generation loop: {e}")
    sys.exit(1)

logging.info("AC Completion Documents have been generated and organized.")
print("AC Completion documents have been generated and organized.")
