import pandas as pd
from docxtpl import DocxTemplate
from docx.shared import Inches
import logging
import os
from datetime import datetime
import time
import argparse
import shutil
import subprocess
import sys
from docx import Document

# ------------------------------
# 1) Parse command-line arguments (including --templates_dir)
# ------------------------------
parser = argparse.ArgumentParser(description='Generate Anger Control Exit Notices')
parser.add_argument('--csv_file', required=True, help='Path to the CSV file')
parser.add_argument('--start_date', required=True, help='Start date in YYYY-MM-DD format')
parser.add_argument('--end_date', required=True, help='End date in YYYY-MM-DD format')
parser.add_argument('--clinic_folder', required=True, help='Clinic folder identifier')

# NEW: templates_dir is required
parser.add_argument('--templates_dir', required=True, help='Path to the clinic-specific templates folder')

parser.add_argument('--output_dir', required=False, help='Optional path to the output directory')
args = parser.parse_args()

csv_file_path   = args.csv_file
start_date_str  = args.start_date
end_date_str    = args.end_date
clinic_folder   = args.clinic_folder
templates_dir   = args.templates_dir   # <--- new

# --- env overrides + templates root ---
program_code  = os.environ.get('NP_PROGRAM_CODE', 'AC')
templates_dir = os.environ.get('NP_TEMPLATES_DIR', templates_dir)
clinic_folder = os.environ.get('NP_CLINIC_FOLDER', clinic_folder)
if os.path.basename(templates_dir) == clinic_folder:
    _TEMPLATES_ROOT = os.path.dirname(templates_dir)
else:
    _TEMPLATES_ROOT = templates_dir

output_dir      = args.output_dir

# ------------------------------
# 2) Logging Setup
# ------------------------------
log_dir = '/home/notesao/NotePro-Report-Generator/logs'
os.makedirs(log_dir, exist_ok=True)
logging.basicConfig(
    filename=os.path.join(log_dir, f'AC_Exit_Notices_{datetime.now().strftime("%Y%m%d")}.log'),
    level=logging.DEBUG,
    format='%(asctime)s [%(levelname)s] %(message)s'
)
logging.info("AC Exit Notices Script started.")

# ------------------------------
# 3) Convert start/end dates
# ------------------------------
try:
    start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
    end_date   = datetime.strptime(end_date_str, '%Y-%m-%d')
except ValueError as e:
    logging.error(f"Invalid date format for start_date or end_date: {e}")
    sys.exit(1)

# ------------------------------
# 4) Define directories
# ------------------------------
base_dir = os.path.dirname(os.path.abspath(__file__))
today_str = datetime.now().strftime('%m.%d.%y')

if output_dir:
    generated_documents_dir = output_dir
else:
    generated_documents_dir = os.path.join(base_dir, 'GeneratedDocuments', clinic_folder, today_str)
os.makedirs(generated_documents_dir, exist_ok=True)

public_generated_documents_dir = f"/home/notesao/{clinic_folder}/public_html/GeneratedDocuments"
os.makedirs(public_generated_documents_dir, exist_ok=True)

central_location_path = os.path.join(public_generated_documents_dir, today_str)
os.makedirs(central_location_path, exist_ok=True)
logging.info(f"Central location directory created: {central_location_path}")

errors_txt = os.path.join(generated_documents_dir, 'errors.txt')
def record_error(msg: str):
    try:
        with open(errors_txt, 'a', encoding='utf-8') as f:
            f.write(str(msg).rstrip() + '\n')
    except Exception:
        pass
    logging.error(msg)

def resolve_template(templates_root, clinic, program, filename):
    for p in (
        os.path.join(templates_root, clinic, filename),
        os.path.join(templates_root, 'default', program, filename),
        os.path.join(templates_root, 'default', 'BIPP', filename),
    ):
        if os.path.isfile(p):
            return p
    return None


# ------------------------------
# 5) Template paths from --templates_dir
# ------------------------------
exit_notice_template_path       = resolve_template(_TEMPLATES_ROOT, clinic_folder, program_code, 'Template.AC Exit Notice.docx')
progress_report_template_path   = resolve_template(_TEMPLATES_ROOT, clinic_folder, program_code, 'Template.AC Exit Progress Report.docx')

for _name, _path in [
    ('Template.AC Exit Notice.docx', exit_notice_template_path),
    ('Template.AC Exit Progress Report.docx', progress_report_template_path),
]:
    if not _path:
        record_error(f"[FATAL] Missing template {_name} under clinic/default fallbacks in {_TEMPLATES_ROOT}.")

# Columns to convert from Y/N → Yes/No
columns_to_convert = [
    "speaks_significantly_in_group", "respectful_to_group",
    "takes_responsibility_for_past", "disruptive_argumentitive",
    "humor_inappropriate", "blames_victim",
    "appears_drug_alcohol", "inappropriate_to_staff"
]

# ------------------------------
# 6) Helper Functions
# ------------------------------
def convert_yn_to_yes_no(value):
    if isinstance(value, str):
        val = value.strip().lower()
        if val == 'y':
            return 'Yes'
        elif val == 'n':
            return 'No'
    return value

def parse_multiple_formats(value, formats=['%m/%d/%Y', '%Y-%m-%d']):
    """Try multiple date formats and return a datetime obj or None."""
    if not value or not isinstance(value, str):
        return None
    value = value.strip()
    for fmt in formats:
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None

def fill_gender_placeholders(context):
    """Fill gender placeholders (gender1..gender6) based on male/female."""
    gender = context.get('gender', '').lower()
    if gender == 'male':
        context['gender1'] = 'his'
        context['gender2'] = 'he'
        context['gender3'] = "Men's"
        context['gender4'] = 'him'
        context['gender5'] = "He"
        context['gender6'] = 'himself'
    elif gender == 'female':
        context['gender1'] = 'her'
        context['gender2'] = 'she'
        context['gender3'] = "Women's"
        context['gender4'] = 'her'
        context['gender5'] = "She"
        context['gender6'] = 'herself'
    else:
        logging.error(f"Unrecognized gender: {gender}. Setting placeholders empty.")
        for g in ['gender1','gender2','gender3','gender4','gender5','gender6']:
            context[g] = ''

def docx_to_pdf(docx_path, pdf_dir, retries=3):
    """Convert DOCX to PDF using LibreOffice with a few retries."""
    attempt = 0
    while attempt < retries:
        try:
            output_file = os.path.join(pdf_dir, os.path.basename(docx_path).replace('.docx', '.pdf'))
            subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', pdf_dir, docx_path],
                           check=True)
            if os.path.exists(output_file):
                logging.info(f"Converted {docx_path} to PDF successfully on attempt {attempt + 1}.")
                return output_file
            else:
                raise FileNotFoundError(f"PDF not found at expected path: {output_file}")
        except subprocess.CalledProcessError as e:
            logging.error(f"Error converting {docx_path} to PDF on attempt {attempt + 1}: {e}")
            attempt += 1
            time.sleep(2)
    logging.error(f"Failed to convert {docx_path} to PDF after {retries} attempts.")
    return None

def sanitize_filename(filename):
    invalid_chars = '<>:"/\\|?*'
    for char in invalid_chars:
        filename = filename.replace(char, '')
    return filename.strip()

def format_m_d_yyyy(value):
    """
    Parses a date string with multiple possible formats
    and returns a string in mm/dd/yyyy format (with leading zeros).
    """
    date_obj = parse_multiple_formats(value, formats=['%m/%d/%Y', '%Y-%m-%d'])
    if date_obj:
        return date_obj.strftime('%m/%d/%Y')
    return value

def fill_placeholders(row):
    """Main placeholder logic: attendance checkmarks, P/A fields, fee, etc."""
    # Mark D1..D30
    for i in range(1, 31):
        row[f'D{i}'] = ''
    row['D1'] = '\u2714'

    attended = row.get('attended', '0')
    required = row.get('required_sessions', '0')
    try:
        attended_f = float(attended)
        required_f = float(required)
    except ValueError:
        attended_f, required_f = 0, 0

    if required_f > 0:
        filled = int((attended_f / required_f) * 30)
        for i in range(2, filled + 1):
            row[f'D{i}'] = '\u2714'

    # P1..P27 / A1..A18 -> short dates mm/dd/yy
    def format_p_a_date(val):
        date_obj = parse_multiple_formats(val)
        if date_obj:
            return date_obj.strftime('%m/%d/%y')
        return val

    for i in range(1, 28):
        p_field = f'P{i}'
        if p_field in row and row[p_field].strip():
            row[p_field] = format_p_a_date(row[p_field])

    for i in range(1, 19):
        a_field = f'A{i}'
        if a_field in row and row[a_field].strip():
            row[a_field] = format_p_a_date(row[a_field])

    if 'client_note' in row and isinstance(row['client_note'], str):
        row['client_note'] = row['client_note'].replace('&', 'and')

    if 'balance' in row:
        try:
            balance = float(row['balance'])
            row['feeproblems'] = "No" if balance == 0 else "Yes"
        except ValueError:
            row['feeproblems'] = "Yes"

    # If last_attended is missing, fallback to orientation_date
    if 'orientation_date' in row and 'last_attended' in row:
        if not row['last_attended'].strip() or row['last_attended'].lower() == 'nan':
            row['last_attended'] = row['orientation_date']

    # Reformat certain date fields to mm/dd/yyyy
    for date_field in ['dob', 'orientation_date', 'last_attended', 'exit_date', 'report_date']:
        row[date_field] = format_m_d_yyyy(row.get(date_field, ''))

    # AC attendance checkmarks for R, C, S, O placeholders
    try:
        attended_f = float(row.get('attended', '0'))
        required_f = float(row.get('required_sessions', '0'))
        if required_f > 0:
            # Clear all checkboxes initially
            for section in ['R', 'C', 'S', 'O']:
                for i in range(1, 4):
                    row[f'{section}{i}'] = ''

            # 1) 1–4  => R1, C1, O1, S1
            if 1 <= attended_f <= 4:
                row['R1'] = row['C1'] = row['O1'] = row['S1'] = '\u2714'

            # 2) 5–13 => R2, C2, O1, S2
            elif 5 <= attended_f <= 13:
                row['R2'] = row['C2'] = row['O1'] = row['S2'] = '\u2714'

            # 3) 14–15 => R2, C2, O1, S2
            elif 14 <= attended_f <= 15:
                row['R2'] = row['C2'] = row['O1'] = '\u2714'
                row['S2'] = '\u2714'  # S2 also checked

            # 4) 16–18 => R2, C2, O1, S2
            elif 16 <= attended_f <= 18:
                row['R2'] = row['C2'] = row['O1'] = row['S2'] = '\u2714'

    except ValueError:
        # If invalid attended/required, just skip
        for section in ['R','C','S','O']:
            for i in range(1,4):
                row[f'{section}{i}'] = ''

    return row

def render_and_save(template_path, doc_type, context):
    if not template_path:
        return None
    """Render the template, convert to PDF, remove the .docx, return PDF path."""
    try:
        doc = DocxTemplate(template_path)
        doc.render(context)

        first_name = sanitize_filename(context.get('first_name', ''))
        last_name = sanitize_filename(context.get('last_name', ''))
        docx_filename = f"{last_name} {first_name} {doc_type}.docx"
        docx_path = os.path.join(generated_documents_dir, docx_filename)

        doc.save(docx_path)
        pdf_path = docx_to_pdf(docx_path, generated_documents_dir)
        if pdf_path:
            os.remove(docx_path)  # remove DOCX after conversion
            return pdf_path
        else:
            logging.error(f"PDF conversion failed for {docx_path}")
            return None
    except Exception as e:
        logging.error(f"Error rendering {doc_type}: {e}")
        return None

def generate_documents_for_row(row):
    if not exit_notice_template_path and not progress_report_template_path:
        return
    context = row.to_dict()
    fill_gender_placeholders(context)

    # Replace 'nan' with empty string in the context
    for key, value in context.items():
        if isinstance(value, str) and value.strip().lower() == 'nan':
            context[key] = ''

    # Add report_date as today's date in m/d/yyyy
    now = datetime.now()
    context['report_date'] = f"{now.month}/{now.day}/{now.year}"

    # 1) Generate AC Exit Notice
    exit_pdf = render_and_save(exit_notice_template_path, "AC Exit Notice", context) if exit_notice_template_path else None
    # 2) Generate AC Exit Progress Report
    progress_pdf = render_and_save(progress_report_template_path, "AC Exit PReport", context) if progress_report_template_path else None

    # --------------------------------
    # Office / Manager Subfolder Logic
    # --------------------------------
    case_manager_office = str(context.get('case_manager_office', '')).strip()
    case_manager_first  = str(context.get('case_manager_first_name', '')).strip()
    case_manager_last   = str(context.get('case_manager_last_name', '')).strip()

    # If office is missing or "nan", place PDFs directly in central_location_path
    if not case_manager_office or case_manager_office.lower() == 'nan':
        office_dir = central_location_path
    else:
        office_dir = os.path.join(central_location_path, sanitize_filename(case_manager_office))

    # If first name is missing/'nan', only use last name
    if not case_manager_first or case_manager_first.lower() == 'nan':
        manager_subfolder_name = sanitize_filename(case_manager_last)
    else:
        manager_subfolder_name = f"{sanitize_filename(case_manager_first)} {sanitize_filename(case_manager_last)}"

    manager_subfolder_name = manager_subfolder_name.strip()
    manager_dir = os.path.join(office_dir, manager_subfolder_name)
    os.makedirs(manager_dir, exist_ok=True)

    # Move PDFs
    for pdf_file in [exit_pdf, progress_pdf]:
        if pdf_file and os.path.exists(pdf_file):
            destination_path = os.path.join(manager_dir, os.path.basename(pdf_file))
            shutil.move(pdf_file, destination_path)
            logging.info(f"Moved {pdf_file} to {destination_path}")
        else:
            logging.error("PDF conversion failed or file doesn't exist; document not moved.")

# ------------------------------
# 7) Load & Filter CSV
# ------------------------------
try:
    df = pd.read_csv(csv_file_path)
    logging.info(f"Loaded {len(df)} records from CSV: {csv_file_path}")
except Exception as e:
    logging.error(f"Failed to load CSV: {e}")
    sys.exit(1)

# Convert exit_date to datetime
df['exit_date'] = pd.to_datetime(df['exit_date'], errors='coerce')
df = df.dropna(subset=['exit_date'])

# Filter: program_name = 'Anger Control', exit_reason in ['Violation of Requirements','Unable to Participate','Death','Moved']
df = df[df['program_name'] == 'Anger Control']
df = df[df['exit_reason'].isin(['Violation of Requirements', 'Unable to Participate', 'Death', 'Moved'])]

# Filter by date range
df = df[(df['exit_date'] >= start_date) & (df['exit_date'] <= end_date)]

if df.empty:
    logging.warning("No records found after filtering. No documents will be generated.")
    sys.exit(0)

# Convert everything to str for placeholders
df = df.astype(str)

# Fill placeholders
df = df.apply(fill_placeholders, axis=1)
df.fillna('', inplace=True)

# Convert Y/N columns to Yes/No
for col in columns_to_convert:
    if col in df.columns:
        df[col] = df[col].apply(convert_yn_to_yes_no)

# ------------------------------
# 8) Main Loop: generate documents
# ------------------------------
try:
    for _, row in df.iterrows():
        generate_documents_for_row(row)
except Exception as e:
    logging.error(f"Error during document generation loop: {e}")
    sys.exit(1)

logging.info("AC Exit Notices have been generated and organized.")
print("AC Exit Notices have been generated and organized.")
