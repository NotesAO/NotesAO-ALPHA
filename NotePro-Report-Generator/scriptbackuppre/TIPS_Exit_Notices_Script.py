#!/usr/bin/env python3
# TIPS_Exit_Notices_and_Exit_Progress_Reports.py

import pandas as pd
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Inches
import logging
import requests
from io import BytesIO
import os
import sys
from datetime import datetime
from PIL import Image
import time
import argparse
import shutil
import subprocess
import re

##############################################################################
# 1) Parse arguments
##############################################################################
parser = argparse.ArgumentParser(description='Generate TIPS Exit Notices + Exit Progress Reports')
parser.add_argument('--csv_file', required=True, help='Path to the CSV file')
parser.add_argument('--start_date', required=True, help='Start date in YYYY-MM-DD format')
parser.add_argument('--end_date', required=True, help='End date in YYYY-MM-DD format')
parser.add_argument('--clinic_folder', required=True, help='Clinic folder identifier')
parser.add_argument('--output_dir', required=False, help='Optional path to the output directory')
parser.add_argument('--templates_dir', required=True, help='Path to the templates directory')
args = parser.parse_args()

csv_file_path = args.csv_file
start_date_str = args.start_date
end_date_str = args.end_date
clinic_folder = args.clinic_folder
output_dir = args.output_dir
templates_dir = args.templates_dir

##############################################################################
# 2) Set up logging
##############################################################################
log_dir = '/home/notesao/NotePro-Report-Generator/logs'
os.makedirs(log_dir, exist_ok=True)
logging.basicConfig(
    filename=os.path.join(log_dir, f'TIPS_Exit_Notices_{datetime.now().strftime("%Y%m%d")}.log'),
    level=logging.DEBUG,
    format='%(asctime)s [%(levelname)s] %(message)s'
)
logging.info("TIPS Exit Notices + Exit Progress Reports Script started.")

##############################################################################
# 3) Convert date arguments
##############################################################################
try:
    start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
    end_date = datetime.strptime(end_date_str, '%Y-%m-%d')
except ValueError as e:
    logging.error(f"Invalid date format for start_date or end_date: {e}")
    sys.exit(1)

##############################################################################
# 4) Set up directories
##############################################################################
base_dir = os.path.dirname(os.path.abspath(__file__))
today = datetime.now().strftime('%m.%d.%y')

if output_dir:
    generated_documents_dir = output_dir
else:
    generated_documents_dir = os.path.join(base_dir, 'GeneratedDocuments', clinic_folder, today)
os.makedirs(generated_documents_dir, exist_ok=True)

public_generated_documents_dir = f"/home/notesao/{clinic_folder}/public_html/GeneratedDocuments"
os.makedirs(public_generated_documents_dir, exist_ok=True)

central_location_path = os.path.join(public_generated_documents_dir, today)
os.makedirs(central_location_path, exist_ok=True)
logging.info(f"Central location directory created: {central_location_path}")

##############################################################################
# 5) Template paths
##############################################################################
exit_notice_template_path = os.path.join(templates_dir, 'Template.TIPS Exit Notice.docx')
exit_progress_report_template_path = os.path.join(templates_dir, 'Template.TIPS Exit Progress Report.docx')

##############################################################################
# 6) Helper functions
##############################################################################

def sanitize_filename(filename):
    invalid_chars = '<>:"/\\|?*'
    for char in invalid_chars:
        filename = filename.replace(char, '')
    return filename.strip()

def docx_to_pdf(docx_path, pdf_dir, retries=3):
    """Convert DOCX to PDF using LibreOffice with retry logic."""
    attempt = 0
    while attempt < retries:
        try:
            output_file = os.path.join(pdf_dir, os.path.basename(docx_path).replace('.docx', '.pdf'))
            subprocess.run([
                'libreoffice', '--headless',
                '--convert-to', 'pdf',
                '--outdir', pdf_dir,
                docx_path
            ], check=True)
            if os.path.exists(output_file):
                logging.info(f"Converted {docx_path} to PDF successfully on attempt {attempt + 1}.")
                return output_file
            else:
                raise FileNotFoundError(f"PDF not found at expected path: {output_file}")
        except subprocess.CalledProcessError as e:
            logging.error(f"Error converting {docx_path} to PDF on attempt {attempt + 1}: {e}")
            attempt += 1
            time.sleep(2)
    logging.error(f"Failed to convert {docx_path} to PDF after {retries} attempts.")
    return None

def fetch_image(url, doc, max_width=4.79, max_height=3.92, margin=0.2):
    """Fetch and resize an image to insert into the DOCX template."""
    headers = {
        'User-Agent': 'Mozilla/5.0',
        'Accept': 'image/*'
    }
    try:
        response = requests.get(url, headers=headers, timeout=30)
        if response.status_code == 200 and 'image' in response.headers.get('Content-Type', ''):
            img = Image.open(BytesIO(response.content))

            img_width, img_height = img.size
            aspect_ratio = img_width / img_height

            adjusted_max_width = max_width - margin
            adjusted_max_height = max_height - margin

            if aspect_ratio > 1:
                new_width = min(adjusted_max_width, img_width)
                new_height = new_width / aspect_ratio
                if new_height > adjusted_max_height:
                    new_height = adjusted_max_height
                    new_width = new_height * aspect_ratio
            else:
                new_height = min(adjusted_max_height, img_height)
                new_width = new_height * aspect_ratio
                if new_width > adjusted_max_width:
                    new_width = adjusted_max_width
                    new_height = new_width / aspect_ratio

            img = img.resize((int(new_width * 300), int(new_height * 300)), Image.LANCZOS)

            img_byte_arr = BytesIO()
            img.save(img_byte_arr, format='JPEG')

            return InlineImage(doc, BytesIO(img_byte_arr.getvalue()), width=Inches(new_width), height=Inches(new_height))
        else:
            logging.warning("Failed to fetch image or non-image content returned: %s", url)
            return None
    except requests.RequestException as e:
        logging.warning("Request failed for URL %s with error: %s", url, e)
        return None
    except Exception as e:
        logging.error("Unexpected error fetching image: %s", e)
        return None

def parse_multiple_formats(value, formats=None):
    """Try multiple date formats for a single string value."""
    if formats is None:
        # Includes expansions for TIPS Progress logic
        formats = [
            '%m/%d/%Y', '%Y-%m-%d', '%m/%d/%y',
            '%b-%Y', '%B %Y',          # For month-year combos
            '%Y-%m-%d %H:%M:%S'        # Sometimes CSV has timestamp
        ]
    if not value or not isinstance(value, str):
        return None
    for fmt in formats:
        try:
            return datetime.strptime(value.strip(), fmt)
        except ValueError:
            continue
    return None

def convert_yn_to_yes_no(value):
    """Converts a single Y or N to Yes or No, else returns original."""
    if isinstance(value, str):
        val = value.strip().lower()
        if val == 'y':
            return 'Yes'
        elif val == 'n':
            return 'No'
    return value

##############################################################################
# Transform c# columns (like c1_01, c2_14, etc.)
##############################################################################
def transform_calendar_cell(cell_value):
    """
    Transform c#_ columns so that 'P' -> ✅, 'X' -> ❌, 'E' -> ✖️,
    plus decode weird CSV checkmark encodings.
    """
    if not isinstance(cell_value, str) or not cell_value.strip():
        return cell_value  # return as-is if empty or not string

    raw = cell_value.strip()

    # Decode weird sequences or older checkmarks:
    raw = raw.replace("âœ…", "P")
    raw = raw.replace("✅", "P")
    raw = raw.replace("âŒ", "X")
    raw = raw.replace("❌", "X")
    raw = raw.replace("âœ–", "X")
    raw = raw.replace("â–", "E")
    raw = raw.replace("✖", "E")
    # Now you have a uniform set of letters: P, X, E

    # If you want day numbers (like "8 P"), parse them
    parts = raw.split()
    day_number = None
    symbol_string = ""
    for part in parts:
        if part.isdigit():
            day_number = part  # store the day number
        elif part.upper() in ['P', 'X', 'E']:
            symbol_string += part.upper()

    # Mapping from letters to final symbols
    symbol_map = {'P': '✅', 'X': '❌', 'E': '✖️'}
    replaced_symbols = "".join(symbol_map.get(ch, ch) for ch in symbol_string)

    if day_number:
        return f"{day_number} {replaced_symbols}"
    else:
        return replaced_symbols

##############################################################################
# Fill placeholders with the same logic from TIPS Progress script
##############################################################################
def fill_placeholders(row):
    """
    This function merges the standard TIPS progress placeholders logic
    and also works for exit (just ignoring some fields if they're blank).
    """
    # D1..D30 placeholders for attendance squares
    for i in range(1, 31):
        row[f'D{i}'] = ''

    # Calculate how many squares to fill
    attended = row.get('attended', '0')
    required = row.get('required_sessions', '0')
    try:
        attended_f = float(attended)
        required_f = float(required)
    except ValueError:
        attended_f, required_f = 0, 0

    if required_f > 0:
        filled = int((attended_f / required_f) * 30)
        for i in range(1, filled + 1):
            row[f'D{i}'] = '✔'  # or \u2714

    # Format P# and A# fields as needed (mm/dd/yy)
    for i in range(1, 28):
        p_field = f'P{i}'
        if p_field in row and isinstance(row[p_field], str) and row[p_field].strip():
            date_obj = parse_multiple_formats(row[p_field])
            if date_obj:
                row[p_field] = date_obj.strftime('%m/%d/%y')

    for i in range(1, 19):
        a_field = f'A{i}'
        if a_field in row and isinstance(row[a_field], str) and row[a_field].strip():
            date_obj = parse_multiple_formats(row[a_field])
            if date_obj:
                row[a_field] = date_obj.strftime('%m/%d/%y')

    # Convert c#_header fields to "Month Year" if possible
    header_keys = ['c1_header', 'c2_header', 'c3_header', 'c4_header']
    for key in header_keys:
        if key in row and isinstance(row[key], str) and row[key].strip():
            value = row[key].strip()
            date_obj = None
            original_value = value

            # Attempt "MonthNameYYYY" or "Jun-23" etc.
            # 1) Try small known formats
            for fmt in ['%b-%y', '%b %y', '%b/%y']:
                try:
                    date_obj = datetime.strptime(value, fmt)
                    break
                except ValueError:
                    continue

            # 2) "December2024" -> "December 2024"
            if not date_obj:
                match = re.match(r"([A-Za-z]+)(\d{4})", value)
                if match:
                    month_name, year = match.groups()
                    try:
                        date_obj = datetime.strptime(f"{month_name} {year}", "%B %Y")
                    except ValueError:
                        logging.error(f"Failed to parse MonthYear format for {key}: {value}")

            if date_obj:
                formatted_date = date_obj.strftime('%B %Y')
                row[key] = formatted_date
                logging.info(f"Transformed {key}: {original_value} -> {formatted_date}")
            else:
                logging.error(f"Date parsing error for {key}: {value}, leaving as is.")

    # Convert Y/N columns to Yes/No
    # (If you have a known set of columns, you can explicitly convert them or do all)
    for col in row.keys():
        row[col] = convert_yn_to_yes_no(row[col])

    # Replace '&' in client_note
    if 'client_note' in row and isinstance(row['client_note'], str):
        row['client_note'] = row['client_note'].replace('&', 'and')

    # If a 'balance' is present, set feeproblems
    if 'balance' in row:
        try:
            b = float(row['balance'])
            row['feeproblems'] = "No" if b == 0 else "Yes"
        except ValueError:
            row['feeproblems'] = "Yes"

    # Gender placeholders
    gender = row.get('gender', '').lower()
    if gender == 'male':
        row['gender1'] = 'his'
        row['gender2'] = 'he'
        row['gender3'] = "Men's"
        row['gender4'] = 'him'
        row['gender5'] = 'He'
        row['gender6'] = 'himself'
    elif gender == 'female':
        row['gender1'] = 'her'
        row['gender2'] = 'she'
        row['gender3'] = "Women's"
        row['gender4'] = 'her'
        row['gender5'] = 'She'
        row['gender6'] = 'herself'
    else:
        logging.error(f"Unrecognized gender: {gender}")
        for g in ['gender1','gender2','gender3','gender4','gender5','gender6']:
            row[g] = ''

    # If last_attended is missing, fallback to orientation_date if available
    if 'orientation_date' in row and 'last_attended' in row:
        if not row['last_attended'].strip() or row['last_attended'].lower() == 'nan':
            row['last_attended'] = row['orientation_date']

    # Convert "age" to an integer if possible
    if 'age' in row:
        try:
            age_value = float(row['age'])
            if age_value.is_integer():
                row['age'] = str(int(age_value))
            else:
                row['age'] = str(age_value)
        except ValueError:
            pass

    # Explicitly parse & reformat DOB, orientation_date, last_attended => mm/dd/yyyy
    date_fields_to_format = ['dob', 'orientation_date', 'last_attended', 'report_date', 'exit_date']
    for field in date_fields_to_format:
        if field in row and isinstance(row[field], str) and row[field].strip():
            date_obj = parse_multiple_formats(row[field])
            if date_obj:
                row[field] = date_obj.strftime('%m/%d/%Y')

    return row

def render_and_save(template_path, doc_type, context):
    """Render a single docx (template_path) -> PDF, given context."""
    try:
        doc = DocxTemplate(template_path)

        # Insert image if available
        if 'image_url' in context and context['image_url']:
            image = fetch_image(context['image_url'], doc)
            context['image'] = image if image else ''

        logging.info(f"Context before rendering {doc_type}: {context}")
        doc.render(context)

        # Construct docx path
        first_name = sanitize_filename(context.get('first_name', ''))
        last_name = sanitize_filename(context.get('last_name', ''))
        docx_filename = f"{last_name} {first_name} {doc_type}.docx"
        docx_path = os.path.join(generated_documents_dir, docx_filename)

        doc.save(docx_path)

        # Convert to PDF
        pdf_path = docx_to_pdf(docx_path, generated_documents_dir)
        if pdf_path:
            os.remove(docx_path)  # Remove .docx after successful PDF creation
            return pdf_path
        else:
            logging.error(f"PDF conversion failed for {docx_path}")
            return None
    except Exception as e:
        logging.error(f"Error rendering {doc_type}: {e}")
        return None

##############################################################################
# 7) Load and filter CSV
##############################################################################
try:
    df = pd.read_csv(csv_file_path)
    logging.info(f"Loaded {len(df)} records from CSV.")
except Exception as e:
    logging.error(f"Failed to load CSV: {e}")
    sys.exit(1)

# Must be TIPS. Must have exit_date in the date range. Must have exit_reason not in ["Not Exited", "Completion of Program"].
df['exit_date'] = pd.to_datetime(df['exit_date'], errors='coerce')
df = df.dropna(subset=['exit_date'])  # drop if exit_date can't parse

df = df[df['program_name'] == 'TIPS (Theft)']
df = df[~df['exit_reason'].isin(['Not Exited', 'Completion of Program'])]

df = df[(df['exit_date'] >= start_date) & (df['exit_date'] <= end_date)]

if df.empty:
    logging.warning("No records found after filtering. No documents will be generated.")
    sys.exit(0)

##############################################################################
# 8) Prepare columns and placeholders
##############################################################################
# Convert all columns to string
df = df.astype(str)
df.fillna('', inplace=True)

# Apply fill_placeholders (progress‐style logic)
df = df.apply(fill_placeholders, axis=1)
df.replace(['nan', 'NaN'], '', inplace=True)

# Identify day‐by‐day calendar columns
calendar_cols = []
for c in df.columns:
    if c.startswith(('c1_', 'c2_', 'c3_', 'c4_')) and not c.endswith('_header'):
        calendar_cols.append(c)

# Transform the day cells
for col in calendar_cols:
    df[col] = df[col].apply(transform_calendar_cell)

##############################################################################
# 9) Generate both documents (Exit Notice & Exit Progress Report)
##############################################################################
def generate_documents_for_row(row):
    context = row.to_dict()

    # Add report_date as today's date m/d/yyyy
    now = datetime.now()
    context['report_date'] = now.strftime('%m/%d/%Y')


    # 1) Render TIPS Exit Notice
    exit_pdf = render_and_save(exit_notice_template_path, "TIPS Exit Notice", context)

    # 2) Render TIPS Exit Progress Report
    progress_pdf = render_and_save(exit_progress_report_template_path, "TIPS Exit Progress Report", context)

    # Organize PDFs by case manager
    case_manager_office = sanitize_filename(context.get('case_manager_office', 'Unknown Office'))
    case_manager_first = sanitize_filename(context.get('case_manager_first_name', ''))
    case_manager_last = sanitize_filename(context.get('case_manager_last_name', ''))

    office_dir = os.path.join(central_location_path, case_manager_office)
    manager_dir = os.path.join(office_dir, f"{case_manager_first} {case_manager_last}")
    os.makedirs(manager_dir, exist_ok=True)

    for pdf_file in [exit_pdf, progress_pdf]:
        if pdf_file and os.path.exists(pdf_file):
            new_path = os.path.join(manager_dir, os.path.basename(pdf_file))
            shutil.move(pdf_file, new_path)
            logging.info(f"Moved {pdf_file} to {new_path}")
        else:
            if pdf_file:
                logging.warning(f"PDF file {pdf_file} not generated or missing.")

##############################################################################
# 10) Main loop
##############################################################################
try:
    for _, row in df.iterrows():
        generate_documents_for_row(row)
except Exception as e:
    logging.error(f"Error during document generation loop: {e}")
    sys.exit(1)

logging.info("TIPS Exit Notices and Exit Progress Reports have been generated and organized.")
print("TIPS Exit Notices and Exit Progress Reports have been generated and organized.")
