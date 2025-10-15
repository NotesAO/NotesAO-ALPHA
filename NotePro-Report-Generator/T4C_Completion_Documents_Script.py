#!/usr/bin/env python3
# T4C_Completion_Documents_Script.py

import pandas as pd
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Inches
import logging
import requests
from io import BytesIO
import os
import sys
from datetime import datetime
from PIL import Image
import time
import argparse
import shutil
import subprocess
import re

##############################################################################
# 1) Parse arguments
##############################################################################
parser = argparse.ArgumentParser(description='Generate T4C Completion Documents (Letter, Certificate, Completion Progress Report)')
parser.add_argument('--csv_file', required=True, help='Path to the CSV file')
parser.add_argument('--start_date', required=True, help='Start date in YYYY-MM-DD format')
parser.add_argument('--end_date', required=True, help='End date in YYYY-MM-DD format')
parser.add_argument('--clinic_folder', required=True, help='Clinic folder identifier')
parser.add_argument('--output_dir', required=False, help='Optional path to the output directory')
parser.add_argument('--templates_dir', required=True, help='Path to the templates directory')
args = parser.parse_args()

csv_file_path = args.csv_file
start_date_str = args.start_date
end_date_str = args.end_date
clinic_folder = args.clinic_folder
output_dir = args.output_dir
templates_dir = args.templates_dir

##############################################################################
# 2) Set up logging
##############################################################################
log_dir = '/home/notesao/NotePro-Report-Generator/logs'
os.makedirs(log_dir, exist_ok=True)
logging.basicConfig(
    filename=os.path.join(log_dir, f'T4C_Completion_{datetime.now().strftime("%Y%m%d")}.log'),
    level=logging.DEBUG,
    format='%(asctime)s [%(levelname)s] %(message)s'
)
logging.info("T4C Completion Documents Script started.")

##############################################################################
# 3) Convert date arguments
##############################################################################
try:
    start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
    end_date = datetime.strptime(end_date_str, '%Y-%m-%d')
except ValueError as e:
    logging.error(f"Invalid date format for start_date or end_date: {e}")
    sys.exit(1)

##############################################################################
# 4) Set up directories
##############################################################################
base_dir = os.path.dirname(os.path.abspath(__file__))
today = datetime.now().strftime('%m.%d.%y')

if output_dir:
    generated_documents_dir = output_dir
else:
    generated_documents_dir = os.path.join(base_dir, 'GeneratedDocuments', clinic_folder, today)
os.makedirs(generated_documents_dir, exist_ok=True)

public_generated_documents_dir = f"/home/notesao/{clinic_folder}/public_html/GeneratedDocuments"
os.makedirs(public_generated_documents_dir, exist_ok=True)

central_location_path = os.path.join(public_generated_documents_dir, today)
os.makedirs(central_location_path, exist_ok=True)
logging.info(f"Central location directory created: {central_location_path}")

##############################################################################
# 5) Template paths
##############################################################################
completion_letter_template = os.path.join(templates_dir, 'Template.T4C Completion Letter.docx')
completion_certificate_template = os.path.join(templates_dir, 'Template.T4C Completion Certificate.docx')
completion_report_template = os.path.join(templates_dir, 'Template.T4C Completion Progress Report.docx')

##############################################################################
# 6) Helper Functions
##############################################################################

def sanitize_filename(filename):
    invalid_chars = '<>:"/\\|?*'
    for char in invalid_chars:
        filename = filename.replace(char, '')
    return filename.strip()

def docx_to_pdf(docx_path, pdf_dir, retries=3):
    """Convert DOCX to PDF using LibreOffice with retry logic."""
    attempt = 0
    while attempt < retries:
        try:
            output_file = os.path.join(pdf_dir, os.path.basename(docx_path).replace('.docx', '.pdf'))
            subprocess.run([
                'libreoffice', '--headless',
                '--convert-to', 'pdf',
                '--outdir', pdf_dir,
                docx_path
            ], check=True)
            if os.path.exists(output_file):
                logging.info(f"Converted {docx_path} to PDF successfully on attempt {attempt + 1}.")
                return output_file
            else:
                raise FileNotFoundError(f"PDF not found at expected path: {output_file}")
        except subprocess.CalledProcessError as e:
            logging.error(f"Error converting {docx_path} to PDF on attempt {attempt + 1}: {e}")
            attempt += 1
            time.sleep(2)
    logging.error(f"Failed to convert {docx_path} to PDF after {retries} attempts.")
    return None

def fetch_image(url, doc, max_width=4.79, max_height=3.92, margin=0.2):
    """Fetch and resize an image to insert into the DOCX template."""
    headers = {
        'User-Agent': 'Mozilla/5.0',
        'Accept': 'image/*'
    }
    try:
        response = requests.get(url, headers=headers, timeout=30)
        if response.status_code == 200 and 'image' in response.headers.get('Content-Type', ''):
            img = Image.open(BytesIO(response.content))

            img_width, img_height = img.size
            aspect_ratio = img_width / img_height

            adjusted_max_width = max_width - margin
            adjusted_max_height = max_height - margin

            # Keep aspect ratio, but limit to max_width / max_height
            if aspect_ratio > 1:
                new_width = min(adjusted_max_width, img_width)
                new_height = new_width / aspect_ratio
                if new_height > adjusted_max_height:
                    new_height = adjusted_max_height
                    new_width = new_height * aspect_ratio
            else:
                new_height = min(adjusted_max_height, img_height)
                new_width = new_height * aspect_ratio
                if new_width > adjusted_max_width:
                    new_width = adjusted_max_width
                    new_height = new_width / aspect_ratio

            # Convert from inches to pixels at ~300 DPI
            img = img.resize((int(new_width * 300), int(new_height * 300)), Image.LANCZOS)

            img_byte_arr = BytesIO()
            img.save(img_byte_arr, format='JPEG')

            return InlineImage(doc, BytesIO(img_byte_arr.getvalue()), width=Inches(new_width), height=Inches(new_height))
        else:
            logging.warning("Failed to fetch image or non-image content returned: %s", url)
            return None
    except requests.RequestException as e:
        logging.warning("Request failed for URL %s with error: %s", url, e)
        return None
    except Exception as e:
        logging.error("Unexpected error fetching image: %s", e)
        return None

def parse_multiple_formats(value, formats=None):
    """Try multiple date formats for a single string value (expanded for T4C logic)."""
    if formats is None:
        formats = [
            '%m/%d/%Y', '%Y-%m-%d', '%m/%d/%y',
            '%b-%Y', '%B %Y',
            '%Y-%m-%d %H:%M:%S'
        ]
    if not value or not isinstance(value, str):
        return None
    for fmt in formats:
        try:
            return datetime.strptime(value.strip(), fmt)
        except ValueError:
            continue
    return None

def convert_yn_to_yes_no(value):
    """Converts 'y' to 'Yes', 'n' to 'No', otherwise returns original."""
    if isinstance(value, str):
        val = value.strip().lower()
        if val == 'y':
            return 'Yes'
        elif val == 'n':
            return 'No'
    return value

def transform_calendar_cell(cell_value):
    """
    Preserve multiple marks in a single calendar day for completion docs.
      P -> ✅ (Present)
      X -> ❌ (Unexcused)
      E -> ✖️ (Excused)

    Handles exact mojibake patterns seen in CSV (incl. orphan trailing 'â').
    Returns: "<day> <marks>" (keeps the space, matching current templates).
    """
    if not isinstance(cell_value, str) or not cell_value.strip():
        return cell_value

    raw = cell_value.strip()

    import re
    # 0) Optional: map letters if they appear
    raw = re.sub(r'(?<!\w)P(?!\w)', '✅', raw)
    raw = re.sub(r'(?<!\w)X(?!\w)', '❌', raw)
    raw = re.sub(r'(?<!\w)E(?!\w)', '✖️', raw)

    # ------------------------------------------------------------------
    # 1) TWO-MARK COMBINATIONS FIRST (most specific → least specific)
    #    Orphan-tail cases: first symbol complete, second truncated to 'â'
    # ------------------------------------------------------------------

    # Present + orphan  → Present + Un/Exc (we use Unexcused per your data)
    raw = raw.replace("âœ… â ", "✅❌")
    raw = raw.replace("âœ… â",  "✅❌")
    raw = raw.replace("â\x9c\x85 â", "✅❌")

    # Excused + orphan  → Excused + Excused
    raw = raw.replace("âœ– â ", "✖️✖️")
    raw = raw.replace("âœ– â",  "✖️✖️")
    raw = raw.replace("â\x9d– â", "✖️✖️")

    # Unexcused + orphan → Unexcused + Excused
    raw = raw.replace("â\x9d\x8c â ", "❌✖️")
    raw = raw.replace("â\x9d\x8c â",  "❌✖️")
    raw = raw.replace("âŒ â ",       "❌✖️")
    raw = raw.replace("âŒ â",        "❌✖️")

    # Present + Unexcused (explicit forms)
    raw = raw.replace("âœ… â\x9d", "✅❌")
    raw = raw.replace("âœ… â",  "✅❌")
    raw = raw.replace("â\x9c\x85 â\x9d", "✅❌")

    # Double Present (incl. compact/split)
    raw = raw.replace("âœ… âœ ", "✅✅")
    raw = raw.replace("âœ… âœ",  "✅✅")
    raw = raw.replace("â\x9c\x85 âœ", "✅✅")
    raw = raw.replace("âœ…âœ…",   "✅✅")
    raw = raw.replace("â\x9c\x85 â\x9c", "✅✅")

    # Double Unexcused
    raw = raw.replace("â\x9d\x8c â\x9d", "❌❌")
    raw = raw.replace("âŒ â",        "❌❌")
    raw = raw.replace("â\x9d\x8c â ",   "❌❌")
    raw = raw.replace("â\x9d\x8c â",    "❌❌")
    raw = raw.replace("âŒ â ",         "❌❌")
    raw = raw.replace("âŒ â",          "❌❌")

    # Excused + Unexcused (and compact)
    raw = raw.replace("âœ– â\x9d",      "✖️❌")
    raw = raw.replace("âœ– â",         "✖️❌")
    raw = raw.replace("â\x9d– â\x9d",   "✖️❌")
    raw = raw.replace("âœ–â\x9d\x8c",   "✖️❌")
    raw = raw.replace("â\x9d–â\x9d\x8c","✖️❌")

    # Attendance + Excused and reverse
    raw = raw.replace("âœ… âœ–", "✅✖️")
    raw = raw.replace("âœ… â\x9d–", "✅✖️")
    raw = raw.replace("â\x9c\x85 âœ–", "✅✖️")
    raw = raw.replace("â\x9c\x85 â\x9d–", "✅✖️")
    raw = raw.replace("âœ– âœ…", "✖️✅")
    raw = raw.replace("â\x9d– âœ…", "✖️✅")
    raw = raw.replace("âœ– â\x9c\x85", "✖️✅")
    raw = raw.replace("â\x9d– â\x9c\x85", "✖️✅")
    raw = raw.replace("âœ–âœ…", "✖️✅")
    raw = raw.replace("â\x9d–âœ…", "✖️✅")

    # Unexcused + Excused (and reverse) – remaining explicit forms
    raw = raw.replace("â\x9d\x8c âœ–", "❌✖️")
    raw = raw.replace("â\x9d\x8c â\x9d–", "❌✖️")
    raw = raw.replace("âŒ âœ–", "❌✖️")
    raw = raw.replace("âŒ â\x9d–", "❌✖️")

    # Double Excused (compact)
    raw = raw.replace("âœ–âœ–", "✖️✖️")
    raw = raw.replace("â\x9d–â\x9d–", "✖️✖️")

    # -------------------------------------------------------------
    # 2) SINGLE-SYMBOL NORMALIZATION (exact sequences; ABSENCES FIRST)
    # -------------------------------------------------------------
    # Unexcused ❌ (U+274C)
    raw = raw.replace("â\x9d\x8c", "❌")
    raw = raw.replace("âŒ",       "❌")

    # Excused ✖️ (U+2716 + VS16)
    raw = raw.replace("âœ–",       "✖️")
    raw = raw.replace("â\x9d–",     "✖️")
    raw = raw.replace("✖",         "✖️")  # normalize to VS16

    # Present ✅ (U+2705)
    raw = raw.replace("âœ…",       "✅")
    raw = raw.replace("â\x9c\x85", "✅")

    # -------------------------------------------------------------
    # 3) CLEANUP: orphan mojibake only when adjacent to a real symbol
    # -------------------------------------------------------------
    raw = re.sub(r'[Ââ]\s*(?=(✅|❌|✖️))', '', raw)     # before symbol
    raw = re.sub(r'(✅|❌|✖️)\s*[Ââ]+', r'\1', raw)     # after symbol
    raw = re.sub(r'\s*(✅|❌|✖️)\s*', r'\1', raw)       # tighten spaces

    # -------------------------------------------------------------
    # 4) Extract: first day number + all symbols in encounter order
    # -------------------------------------------------------------
    m = re.search(r'\b(\d{1,2})\b', raw)
    day = m.group(1) if m else None
    marks = ''.join(re.findall(r'✅|❌|✖️', raw))

    # Keep your completion template’s space between day and marks
    if day:
        out = f"{day} {marks}".strip()
    else:
        out = marks.strip()

    # Final micro-cleanups (display artifacts)
    out = out.replace("Â", "").replace("â", "").replace("œ", "").strip()
    return out


##############################################################################
# fill_placeholders for T4C completion progress logic
##############################################################################
def fill_placeholders(row):
    """
    Merges the standard T4C progress placeholders logic (like attendance squares,
    c#_header transformations, date formatting, etc.) for a "Completion" scenario.
    """

    # Convert 'nan' or 'NaN' to empty
    for k, v in row.items():
        if isinstance(v, str) and v.strip().lower() == 'nan':
            row[k] = ''

    # D1..D30 placeholders
    for i in range(1, 31):
        row[f'D{i}'] = ''

    # Attendance squares
    attended = row.get('attended', '0')
    required = row.get('required_sessions', '0')
    try:
        attended_f = float(attended)
        required_f = float(required)
    except ValueError:
        attended_f, required_f = 0, 0

    if required_f > 0:
        filled = int((attended_f / required_f) * 30)
        for i in range(1, filled + 1):
            row[f'D{i}'] = '✔'

    # P# and A# fields -> mm/dd/yy
    for i in range(1, 28):
        p_field = f'P{i}'
        if p_field in row and isinstance(row[p_field], str) and row[p_field].strip():
            d = parse_multiple_formats(row[p_field])
            if d:
                row[p_field] = d.strftime('%m/%d/%y')

    for i in range(1, 19):
        a_field = f'A{i}'
        if a_field in row and isinstance(row[a_field], str) and row[a_field].strip():
            d = parse_multiple_formats(row[a_field])
            if d:
                row[a_field] = d.strftime('%m/%d/%y')

    # Convert c#_header -> "Month Year" if possible
    header_keys = ['c1_header','c2_header','c3_header','c4_header']
    for key in header_keys:
        if key in row and isinstance(row[key], str) and row[key].strip():
            original_val = row[key].strip()
            date_obj = None

            # Attempt small known formats (Jun-23, Jun 23, etc.)
            for fmt in ['%b-%y', '%b %y', '%b/%y']:
                try:
                    date_obj = datetime.strptime(original_val, fmt)
                    break
                except ValueError:
                    continue

            # Try "December2024" => "December 2024"
            if not date_obj:
                match = re.match(r"([A-Za-z]+)(\d{4})", original_val)
                if match:
                    month_name, year = match.groups()
                    try:
                        date_obj = datetime.strptime(f"{month_name} {year}", "%B %Y")
                    except ValueError:
                        logging.error(f"Failed parse for {key} = {original_val}")

            if date_obj:
                row[key] = date_obj.strftime('%B %Y')
                logging.info(f"Transformed {key}: {original_val} -> {row[key]}")
            else:
                logging.error(f"Date parsing error for {key}={original_val}, leaving as is.")

    # Convert Y/N fields -> Yes/No
    for col in row.keys():
        row[col] = convert_yn_to_yes_no(row[col])

    # Replace '&' in client_note
    if 'client_note' in row and isinstance(row['client_note'], str):
        row['client_note'] = row['client_note'].replace('&', 'and')

    # Fee problems
    if 'balance' in row:
        try:
            bal = float(row['balance'])
            row['feeproblems'] = "No" if bal == 0 else "Yes"
        except ValueError:
            row['feeproblems'] = "Yes"

    # Gender placeholders
    gender = row.get('gender', '').lower()
    if gender == 'male':
        row['gender1'] = 'his'
        row['gender2'] = 'he'
        row['gender3'] = "Men's"
        row['gender4'] = 'him'
        row['gender5'] = 'He'
        row['gender6'] = 'himself'
    elif gender == 'female':
        row['gender1'] = 'her'
        row['gender2'] = 'she'
        row['gender3'] = "Women's"
        row['gender4'] = 'her'
        row['gender5'] = 'She'
        row['gender6'] = 'herself'
    else:
        logging.error(f"Unrecognized gender: {gender}")
        for g in ['gender1','gender2','gender3','gender4','gender5','gender6']:
            row[g] = ''

    # If last_attended is missing, fallback to orientation_date
    if 'orientation_date' in row and 'last_attended' in row:
        if not row['last_attended'].strip():
            row['last_attended'] = row['orientation_date']

    # Convert "age" to int if possible
    if 'age' in row:
        try:
            age_val = float(row['age'])
            if age_val.is_integer():
                row['age'] = str(int(age_val))
            else:
                row['age'] = str(age_val)
        except ValueError:
            pass

    # Also parse & format specific date fields => mm/dd/yyyy
    date_fields = ['dob','orientation_date','last_attended','exit_date','report_date']
    for dfld in date_fields:
        if dfld in row and isinstance(row[dfld], str) and row[dfld].strip():
            d_obj = parse_multiple_formats(row[dfld])
            if d_obj:
                row[dfld] = d_obj.strftime('%m/%d/%Y')

    return row

def render_and_save(template_path, doc_type, context):
    """Render a single DOCX -> PDF, then remove the DOCX, returning the PDF path."""
    try:
        doc = DocxTemplate(template_path)

        # Insert image if available
        if 'image_url' in context and context['image_url']:
            image = fetch_image(context['image_url'], doc)
            context['image'] = image if image else ''
        else:
            context['image'] = ''

        logging.info(f"Context before rendering {doc_type}: {context}")
        doc.render(context)

        # Construct DOCX path
        first_name = sanitize_filename(context.get('first_name', ''))
        last_name = sanitize_filename(context.get('last_name', ''))
        docx_filename = f"{last_name} {first_name} {doc_type}.docx"
        docx_path = os.path.join(generated_documents_dir, docx_filename)

        doc.save(docx_path)

        # Convert to PDF
        pdf_path = docx_to_pdf(docx_path, generated_documents_dir)
        if pdf_path:
            os.remove(docx_path)  # Remove .docx after successful PDF
            return pdf_path
        else:
            logging.error(f"PDF conversion failed for {docx_path}")
            return None
    except Exception as e:
        logging.error(f"Error rendering {doc_type}: {e}")
        return None

##############################################################################
# 7) Load & Filter CSV
##############################################################################
try:
    df = pd.read_csv(csv_file_path)
    logging.info(f"Loaded {len(df)} records from CSV.")
except Exception as e:
    logging.error(f"Failed to load CSV: {e}")
    sys.exit(1)

# Filter only T4C, exit_reason == 'Completion of Program', exit_date in range
df['exit_date'] = pd.to_datetime(df['exit_date'], errors='coerce')
df = df.dropna(subset=['exit_date'])

df = df[
    (df['program_name'] == 'Thinking for a Change') &
    (df['exit_reason'] == 'Completion of Program') &
    (df['exit_date'] >= start_date) & (df['exit_date'] <= end_date)
]

if df.empty:
    logging.warning("No records found after filtering. No documents will be generated.")
    sys.exit(0)

##############################################################################
# 8) Prepare placeholders
##############################################################################
df = df.astype(str)
df.fillna('', inplace=True)

# 1) Apply fill_placeholders
df = df.apply(fill_placeholders, axis=1)
df.replace(['nan','NaN'], '', inplace=True)

# 2) Identify day-by-day columns c1_XX, c2_XX, c3_XX, c4_XX (not headers)
calendar_cols = []
for c in df.columns:
    if c.startswith(('c1_', 'c2_', 'c3_', 'c4_')) and not c.endswith('_header'):
        calendar_cols.append(c)

# 3) Transform them
for col in calendar_cols:
    df[col] = df[col].apply(transform_calendar_cell)

##############################################################################
# 9) Generate 3 Documents (Letter, Certificate, Completion Progress Report)
##############################################################################
def generate_documents_for_row(row):
    # Make sure 'report_date' is mm/dd/yyyy today
    now = datetime.now()
    row['report_date'] = now.strftime('%m/%d/%Y')

    context = row.to_dict()

    # 1) T4C Completion Progress Report
    pdf_report = render_and_save(completion_report_template, "T4C Completion Progress Report", context)

    # 2) T4C Completion Letter
    pdf_letter = render_and_save(completion_letter_template, "T4C Completion Letter", context)

    # 3) T4C Completion Certificate
    pdf_certificate = render_and_save(completion_certificate_template, "T4C Completion Certificate", context)

    # Organize PDFs by case manager
    case_manager_office = sanitize_filename(context.get('case_manager_office', 'Unknown Office'))
    cm_first = sanitize_filename(context.get('case_manager_first_name', ''))
    cm_last  = sanitize_filename(context.get('case_manager_last_name', ''))

    office_dir = os.path.join(central_location_path, case_manager_office)
    manager_dir = os.path.join(office_dir, f"{cm_first} {cm_last}".strip())
    os.makedirs(manager_dir, exist_ok=True)

    for pdf_file in [pdf_report, pdf_letter, pdf_certificate]:
        if pdf_file and os.path.exists(pdf_file):
            new_path = os.path.join(manager_dir, os.path.basename(pdf_file))
            shutil.move(pdf_file, new_path)
            logging.info(f"Moved {pdf_file} → {new_path}")
        else:
            if pdf_file:
                logging.warning(f"PDF file {pdf_file} not found or not generated.")

##############################################################################
# 10) Main loop
##############################################################################
try:
    for _, row in df.iterrows():
        generate_documents_for_row(row)
except Exception as e:
    logging.error(f"Error during document generation loop: {e}")
    sys.exit(1)

logging.info("T4C Completion Documents have been generated and organized.")
print("T4C Completion Documents have been generated and organized.")
