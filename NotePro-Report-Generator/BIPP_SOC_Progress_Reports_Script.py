import pandas as pd
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Inches
import logging
import requests
from io import BytesIO
import os
import sys
from datetime import datetime
from PIL import Image
import argparse
import shutil
import subprocess
from docx.enum.table import WD_TABLE_ALIGNMENT
import time

parser = argparse.ArgumentParser(description='Generate BIPP Progress Reports')
parser.add_argument('--csv_file', required=True, help='Path to the CSV file')
parser.add_argument('--start_date', required=False, help='(Ignored) Start date for the reports')
parser.add_argument('--end_date', required=False, help='(Ignored) End date for the reports')
parser.add_argument('--output_dir', required=False, help='Path to the output directory')
parser.add_argument('--clinic_folder', required=True, help='Clinic folder identifier')
parser.add_argument('--templates_dir', required=True, help='Path to the templates directory')
args = parser.parse_args()

csv_file_path = args.csv_file
output_dir = args.output_dir
clinic_folder = args.clinic_folder

# --- Env overrides to match other BIPP scripts ---
program_code  = os.environ.get('NP_PROGRAM_CODE', 'BIPP')
templates_dir = os.environ.get('NP_TEMPLATES_DIR', args.templates_dir)
clinic_folder = os.environ.get('NP_CLINIC_FOLDER', clinic_folder)

# DWAG gate
CLINIC_IS_DWAG = str(clinic_folder).strip().lower() == "dwag"


# right after env overrides
if os.path.basename(templates_dir) == clinic_folder:
    _TEMPLATES_ROOT = os.path.dirname(templates_dir)
else:
    _TEMPLATES_ROOT = templates_dir

log_dir = '/home/notesao/NotePro-Report-Generator/logs'
os.makedirs(log_dir, exist_ok=True)
logging.basicConfig(
    filename=os.path.join(log_dir, f'BIPP_Progress_Reports_{datetime.now().strftime("%Y%m%d")}.log'),
    level=logging.DEBUG,
    format='%(asctime)s [%(levelname)s] %(message)s'
)
logging.info("BIPP Progress Reports Script started.")

base_dir = os.path.dirname(os.path.abspath(__file__))

today = datetime.now().strftime('%m.%d.%y')
if output_dir:
    generated_documents_dir = output_dir
else:
    generated_documents_dir = os.path.join(base_dir, 'GeneratedDocuments', clinic_folder, today)
os.makedirs(generated_documents_dir, exist_ok=True)

public_generated_documents_dir = f"/home/notesao/{clinic_folder}/public_html/GeneratedDocuments"
os.makedirs(public_generated_documents_dir, exist_ok=True)

central_location_path = os.path.join(public_generated_documents_dir, today)
os.makedirs(central_location_path, exist_ok=True)
logging.info(f"Central location directory created: {central_location_path}")

# Per-run human-friendly error sink
errors_txt = os.path.join(central_location_path, 'errors.txt')

def record_error(msg: str):
    try:
        with open(errors_txt, 'a', encoding='utf-8') as f:
            f.write(msg.rstrip() + '\n')
    except Exception as e:
        logging.error(f"Failed writing to errors.txt: {e}")
    logging.warning(msg)


def resolve_template(templates_root: str, clinic: str, program: str, filename: str):
    candidates = [
        os.path.join(templates_root, clinic, filename),
        os.path.join(templates_root, 'default', program, filename),
        os.path.join(templates_root, 'default', 'BIPP', filename),
    ]
    for p in candidates:
        if os.path.isfile(p):
            return p
    return None

# in need_template(...)
def need_template(filename):
    p = resolve_template(_TEMPLATES_ROOT, clinic_folder, program_code, filename)
    if not p:
        record_error(f"[FATAL] Missing template {filename} under clinic/default fallbacks in {_TEMPLATES_ROOT}.")
        sys.exit(0)
    return p

progress_report_template_path         = need_template('Template.BIPP SOC Progress Report.docx')
virtual_progress_report_template_path = need_template('Template.BIPP SOC Progress Report Virtual.docx')



def docx_to_pdf(docx_path, pdf_dir, retries=3):
    attempt = 0
    while attempt < retries:
        try:
            output_file = os.path.join(pdf_dir, os.path.basename(docx_path).replace('.docx', '.pdf'))
            subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', pdf_dir, docx_path], check=True)
            if os.path.exists(output_file):
                logging.info(f"Converted {docx_path} to PDF successfully on attempt {attempt + 1}.")
                return output_file
            else:
                raise FileNotFoundError(f"PDF not found at expected path: {output_file}")
        except subprocess.CalledProcessError as e:
            logging.error(f"Error converting {docx_path} to PDF on attempt {attempt + 1}: {e}")
            attempt += 1
            time.sleep(2)
    logging.error(f"Failed to convert {docx_path} to PDF after {retries} attempts.")
    return None

def align_tables_left(doc):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.LEFT

def fetch_image(url, doc, max_width=3.45, max_height=4.71, margin=0.2):
    headers = {
        'User-Agent': 'Mozilla/5.0',
        'Accept': 'image/*'
    }
    try:
        if not isinstance(url, str) or not url.lower().startswith(('http://','https://')):
            logging.warning(f"Skipping non-http image_url: {url}")
            return None
        response = requests.get(url, headers=headers, timeout=30)
        if response.status_code == 200 and 'image' in response.headers.get('Content-Type', ''):
            img = Image.open(BytesIO(response.content))

            img_width, img_height = img.size
            aspect_ratio = img_width / img_height

            adjusted_max_width = max_width - margin
            adjusted_max_height = max_height - margin

            if aspect_ratio > 1:
                new_width = min(adjusted_max_width, img_width)
                new_height = new_width / aspect_ratio
                if new_height > adjusted_max_height:
                    new_height = adjusted_max_height
                    new_width = new_height * aspect_ratio
            else:
                new_height = min(adjusted_max_height, img_height)
                new_width = new_height * aspect_ratio
                if new_width > adjusted_max_width:
                    new_width = adjusted_max_width
                    new_height = new_width / aspect_ratio

            img = img.resize((int(new_width * 300), int(new_height * 300)), Image.LANCZOS)

            img_byte_arr = BytesIO()
            img.save(img_byte_arr, format='JPEG')

            return InlineImage(doc, BytesIO(img_byte_arr.getvalue()), width=Inches(new_width), height=Inches(new_height))
        else:
            logging.warning("Failed to fetch image or non-image content returned: %s", url)
            return None
    except requests.RequestException as e:
        logging.warning("Request failed for URL %s with error: %s", url, e)
        return None
    except Exception as e:
        logging.error("Unexpected error fetching image: %s", e)
        return None

def convert_yn_to_yes_no(value):
    if isinstance(value, str):
        val = value.strip().lower()
        if val == 'y':
            return 'Yes'
        elif val == 'n':
            return 'No'
    return value

def parse_and_format_date(date_str):
    if not date_str or not isinstance(date_str, str):
        return ''
    date_str = date_str.strip()
    if not date_str:
        return ''
    for fmt in ('%Y-%m-%d', '%m/%d/%Y'):
        try:
            d = datetime.strptime(date_str, fmt)
            return f"{d.month}/{d.day}/{d.year}"  # m/d/yyyy
        except ValueError:
            continue
    return date_str

def parse_and_format_p_a_date(date_str):
    # Format as mm/dd/yy
    if not date_str or not isinstance(date_str, str):
        return ''
    date_str = date_str.strip()
    if not date_str:
        return ''
    for fmt in ('%Y-%m-%d', '%m/%d/%Y'):
        try:
            d = datetime.strptime(date_str, fmt)
            return d.strftime('%m/%d/%y')  # mm/dd/yy
        except ValueError:
            continue
    return ''

# Function to determine fee problems
def determine_fee_problems(balance):
    try:
        balance = float(balance)
        if balance == 0:
            return "No"
        return "Yes"
    except ValueError:
        return ""

# Ensure balance formatting in placeholders
def format_balance(balance):
    try:
        balance = float(balance)
        return f"${int(balance)}"  # Convert to whole number and format as currency
    except ValueError:
        return "$0"  # Default to $0 if invalid

def determine_stage_of_change(attended_str, required_str):
    try:
        attended = float(attended_str)
        required = float(required_str)
    except ValueError:
        logging.error(f"Invalid attended/required values: attended={attended_str}, required={required_str}")
        return 'Alert'

    if required == 0:
        return 'Alert'

    # Exact logic for 18-week program
    if required == 18:
        if 1 <= attended <= 4:
            return 'Precontemplation'
        elif 5 <= attended <= 8:
            return 'Contemplation'
        elif 9 <= attended <= 11:
            return 'Preparation'
        elif 12 <= attended <= 15:
            return 'Action'
        elif 16 <= attended <= 18:
            return 'Maintenance'
        else:
            logging.warning(f"Unexpected attendance value for BIPP 18-week: attended={attended}")
            return 'Alert'

    # Exact logic for 27-week program
    if required == 27:
        if 1 <= attended <= 7:
            return 'Precontemplation'
        elif 8 <= attended <= 14:
            return 'Contemplation'
        elif 15 <= attended <= 18:
            return 'Preparation'
        elif 19 <= attended <= 23:
            return 'Action'
        elif 24 <= attended <= 27:
            return 'Maintenance'
        else:
            logging.warning(f"Unexpected attendance value for BIPP 27-week: attended={attended}")
            return 'Alert'

    # Adaptive logic for other durations using 18-week proportions
    ratio = attended / required
    if ratio > 0 and ratio <= (4/18):
        return 'Precontemplation'
    elif ratio <= (8/18):
        return 'Contemplation'
    elif ratio <= (11/18):
        return 'Preparation'
    elif ratio <= (15/18):
        return 'Action'
    elif ratio <= 1.0:
        return 'Maintenance'
    else:
        logging.warning(f"Ratio out of expected range for required={required}, attended={attended}")
        return 'Alert'

def _is_yes(v):
    return str(v).strip().lower() in ("yes", "y", "true", "1")

def _chk(flag):
    return "✔" if flag else "✘"


def fill_placeholders(row):
    # Set gender placeholders
    gender = row.get('gender', '').lower()
    if gender == 'male':
        row['gender1'] = 'his'
        row['gender2'] = 'he'
        row['gender3'] = "Men's"
        row['gender4'] = 'him'
        row['gender5'] = "He"
        row['gender6'] = "His"
        row['gender7'] = 'himself'
    elif gender == 'female':
        row['gender1'] = 'her'
        row['gender2'] = 'she'
        row['gender3'] = "Women's"
        row['gender4'] = 'her'
        row['gender5'] = "She"
        row['gender6'] = "Her"
        row['gender7'] = 'herself'
    else:
        # neutral fallbacks
        row['gender1'] = 'their'
        row['gender2'] = 'they'
        row['gender3'] = "Participants"
        row['gender4'] = 'them'
        row['gender5'] = "They"
        row['gender6'] = "Their"
        row['gender7'] = 'themself'


    # Format key dates
    for date_field in ['dob', 'orientation_date', 'last_attended']:
        if date_field in row:
            row[date_field] = parse_and_format_date(row[date_field])

    # Attendance placeholders (D1-D30)
    attended = row.get('attended', '0')
    required = row.get('required_sessions', '0')

    try:
        attended_f = float(attended)
        required_f = float(required)
    except ValueError:
        attended_f, required_f = 0, 0

    # Initialize all placeholders as empty
    for i in range(1, 31):
        row[f'D{i}'] = ''

    # Always fill D1 with a check mark
    row['D1'] = '\u2714'

    if required_f > 0:
        # Calculate how many placeholders to fill based on attendance ratio
        filled = int((attended_f / required_f) * 30)
        # Ensure at least D1 is checked
        for i in range(2, filled + 1):  # Start from D2 since D1 is already filled
            row[f'D{i}'] = '\u2714'

    # Replace '&' with 'and'
    if 'client_note' in row:
        row['client_note'] = row['client_note'].replace('&', 'and')

    # Format P1-P27 and A1-A18
    for i in range(1, 28):
        p_field = f'P{i}'
        if p_field in row:
            row[p_field] = parse_and_format_p_a_date(row[p_field])

    for i in range(1, 19):
        a_field = f'A{i}'
        if a_field in row:
            row[a_field] = parse_and_format_p_a_date(row[a_field])

    # Determine client stage of change
    row['client_stagechange'] = determine_stage_of_change(attended, required)

     # ---------- DWAG-only derived fields ----------
    if CLINIC_IS_DWAG:
        # absences
        try:
            absences = int(float(row.get("absence_unexcused", 0) or 0))
        except Exception:
            absences = 0

        # Attendance Benchmarks
        row["attended_initial_assessment"] = _chk(bool(str(row.get("orientation_date","")).strip()))
        row["consistent_attendance"]      = _chk(absences == 0)

        if absences >= 3:
            row["three_or_more_absences"] = f"{absences} absences"
        elif absences == 0:
            row["three_or_more_absences"] = "No"
        else:
            row["three_or_more_absences"] = "Within 3"

        # Pull participation fields as Yes/No
        respectful   = _is_yes(row.get("respectful_to_group", ""))
        takes_resp   = _is_yes(row.get("takes_responsibility_for_past", ""))
        disruptive   = _is_yes(row.get("disruptive_argumentitive", ""))
        humor_bad    = _is_yes(row.get("humor_inappropriate", ""))
        blames       = _is_yes(row.get("blames_victim", ""))
        appears_ua   = _is_yes(row.get("appears_drug_alcohol", ""))
        staff_issue  = _is_yes(row.get("inappropriate_to_staff", ""))

        # Compliance with Program Parameters
        followed_rules = not (disruptive or humor_bad or appears_ua or staff_issue)
        adhered_agree  = followed_rules and not blames
        participated   = respectful and not disruptive

        row["followed_group_rules"]           = _chk(followed_rules)
        row["adhered_to_participant_agreement"]= _chk(adhered_agree)
        row["participated_in_dialogue"]       = _chk(participated)

        # Accountability
        row["accountability_consistent"] = _chk(takes_resp and not blames)

    return row

def sanitize_filename(filename):
    invalid_chars = '<>:"/\\|?*'
    for char in invalid_chars:
        filename = filename.replace(char, '')
    return filename.strip()

try:
    # Force *all* columns to be read as string
    df = pd.read_csv(csv_file_path, dtype=str)
    logging.info(f"Loaded {len(df)} records from CSV.")
except Exception as e:
    logging.error(f"Failed to load CSV file: {e}")
    sys.exit(1)

# Now fill missing with empty string. No numeric columns exist, so no conflict.
df.fillna('', inplace=True)

# No need for df = df.astype(str) -- we already forced everything to be string.


# Convert Y/N to Yes/No
for col in df.columns:
    df[col] = df[col].apply(convert_yn_to_yes_no)

# Filter for BIPP (male/female) and Not Exited clients
df['exit_date'] = pd.to_datetime(df['exit_date'], errors='coerce')

df = df[df['program_name'].isin(['BIPP (male)', 'BIPP (female)'])]
df = df[df['exit_reason'] == 'Not Exited']

if df.empty:
    logging.warning("No records found for BIPP Not Exited clients. No documents will be generated.")
    sys.exit(0)

# Convert exit_date back to string if needed
df['exit_date'] = df['exit_date'].dt.strftime('%Y-%m-%d').fillna('')

# Handle last_attended
def fix_last_attended(row):
    if not row['last_attended'].strip():
        return row['orientation_date']
    return row['last_attended']

df['last_attended'] = df.apply(fix_last_attended, axis=1)

# Fee problems
# Apply determine_fee_problems to calculate fee problems
df['feeproblems'] = df['balance'].apply(determine_fee_problems)

# Apply format_balance to format the balance as whole numbers
df['balance'] = df['balance'].apply(format_balance)

# Convert attended and required_sessions to numeric if possible
df['attended'] = pd.to_numeric(df['attended'], errors='coerce').fillna(0)
df['required_sessions'] = pd.to_numeric(df['required_sessions'], errors='coerce').fillna(0)

# Apply placeholders
df = df.apply(fill_placeholders, axis=1)

def render_and_save(template_path, doc_type, context, output_dir):
    try:
        doc = DocxTemplate(template_path)
        # Insert image if available
        if 'image_url' in context and context['image_url']:
            image = fetch_image(context['image_url'], doc)
            context['image'] = image if image else ''

        doc.render(context)
        align_tables_left(doc)

        first_name = sanitize_filename(context.get('first_name', ''))
        last_name = sanitize_filename(context.get('last_name', ''))
        docx_filename = f"{last_name} {first_name} BIPP Stage PReport.docx"
        docx_path = os.path.join(output_dir, docx_filename)

        doc.save(docx_path)
        time.sleep(1)

        pdf_path = docx_to_pdf(docx_path, output_dir)
        if pdf_path:
            os.remove(docx_path)
            return pdf_path
        else:
            logging.error(f"PDF conversion failed for {docx_path}")
            return None
    except Exception as e:
        logging.error(f"Error rendering {doc_type}: {e}")
        record_error(f"Render failed for {doc_type} / {context.get('last_name','')}, {context.get('first_name','')}: {e}")
        return None


def generate_documents(row, output_dir_path):
    context = row.to_dict()

    # report_date as today's date in m/d/yyyy
    now = datetime.now()
    context['report_date'] = f"{now.month}/{now.day}/{now.year}"

    # ⬇️ Put this here (replaces your existing group_name / report_template logic)
    group_name = str(context.get("group_name", "")).lower()
    report_template = progress_report_template_path
    if "virtual" in group_name or "women" in group_name:
        report_template = virtual_progress_report_template_path

    pdf_path = render_and_save(report_template, "PReport", context, output_dir_path)

    # Organize PDFs by case manager / facilitator / instructor (clinic differences)
    def pick(d, *keys, default=''):
        for k in keys:
            v = d.get(k, '')
            if isinstance(v, str) and v.strip():
                return v.strip()
        return default

    raw_office = pick(row, 'case_manager_office','facilitator_office','instructor_office', default='')
    office_dir = (output_dir_path if not raw_office or raw_office.lower() in ('nan','none','null','')
                else os.path.join(output_dir_path, sanitize_filename(raw_office)))

    raw_first = pick(row, 'case_manager_first_name','facilitator_first_name','instructor_first_name', default='')
    raw_last  = pick(row, 'case_manager_last_name', 'facilitator_last_name', 'instructor_last_name',  default='')
    manager_subfolder = (f"{sanitize_filename(raw_first)} {sanitize_filename(raw_last)}".strip()
                        or (sanitize_filename(raw_last) if raw_last else 'Unknown Manager'))

    manager_dir = os.path.join(office_dir, manager_subfolder)
    os.makedirs(manager_dir, exist_ok=True)

    if pdf_path and os.path.exists(pdf_path):
        shutil.move(pdf_path, os.path.join(manager_dir, os.path.basename(pdf_path)))
        logging.info(f"Moved {pdf_path} to {manager_dir}")
    else:
        logging.warning("PDF file not found or not generated.")


processed = skipped = 0
for idx, row in df.iterrows():
    try:
        generate_documents(row, central_location_path)
        processed += 1
    except Exception as e:
        record_error(f"[row {idx}] {row.get('last_name','')}, {row.get('first_name','')} :: {e}")
        skipped += 1

logging.info(f"Done. processed={processed} skipped={skipped}")
print(f"processed={processed} skipped={skipped}")
sys.exit(0)


