import pandas as pd
from docxtpl import DocxTemplate
import logging
import os
from datetime import datetime
import subprocess
import time
import argparse
import shutil
import sys
from docx import Document

# ------------------------------
# 1) Parse command-line arguments (including --templates_dir)
# ------------------------------
parser = argparse.ArgumentParser(description='Generate Anger Control Entrance Notifications')
parser.add_argument('--csv_file', required=True, help='Path to the CSV file')
parser.add_argument('--start_date', required=True, help='Start date in YYYY-MM-DD format')
parser.add_argument('--end_date', required=True, help='End date in YYYY-MM-DD format')
parser.add_argument('--clinic_folder', required=True, help='Clinic folder identifier')

# NEW: templates_dir argument
parser.add_argument('--templates_dir', required=True, help='Path to the clinic-specific templates folder')

parser.add_argument('--output_dir', required=False, help='Optional path to the output directory')
args = parser.parse_args()

csv_file_path   = args.csv_file
start_date_str  = args.start_date
end_date_str    = args.end_date
clinic_folder   = args.clinic_folder
templates_dir   = args.templates_dir   # <--- new

# --- env overrides + templates root ---
program_code  = os.environ.get('NP_PROGRAM_CODE', 'AC')
templates_dir = os.environ.get('NP_TEMPLATES_DIR', templates_dir)
clinic_folder = os.environ.get('NP_CLINIC_FOLDER', clinic_folder)
if os.path.basename(templates_dir) == clinic_folder:
    _TEMPLATES_ROOT = os.path.dirname(templates_dir)
else:
    _TEMPLATES_ROOT = templates_dir

output_dir      = args.output_dir

# ------------------------------
# 2) Setup logging
# ------------------------------
log_dir = '/home/notesao/NotePro-Report-Generator/logs'
os.makedirs(log_dir, exist_ok=True)
logging.basicConfig(
    filename=os.path.join(log_dir, f'AC_Entrance_Notifications_{datetime.now().strftime("%Y%m%d")}.log'),
    level=logging.DEBUG,
    format='%(asctime)s [%(levelname)s] %(message)s'
)
logging.info("AC Entrance Notifications Script started.")

# ------------------------------
# 3) Convert start/end dates
# ------------------------------
try:
    start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
    end_date   = datetime.strptime(end_date_str, '%Y-%m-%d')
except ValueError as e:
    logging.error(f"Invalid date format for start_date or end_date: {e}")
    sys.exit(1)

# ------------------------------
# 4) Define directories
# ------------------------------
base_dir = os.path.dirname(os.path.abspath(__file__))
today_str = datetime.now().strftime('%m.%d.%y')

if output_dir:
    generated_documents_dir = output_dir
else:
    generated_documents_dir = os.path.join(base_dir, 'GeneratedDocuments', clinic_folder, today_str)
os.makedirs(generated_documents_dir, exist_ok=True)

public_generated_documents_dir = f"/home/notesao/{clinic_folder}/public_html/GeneratedDocuments"
os.makedirs(public_generated_documents_dir, exist_ok=True)

central_location_path = os.path.join(public_generated_documents_dir, today_str)
os.makedirs(central_location_path, exist_ok=True)
logging.info(f"Central location directory created: {central_location_path}")

errors_txt = os.path.join(generated_documents_dir, 'errors.txt')
def record_error(msg: str):
    try:
        with open(errors_txt, 'a', encoding='utf-8') as f:
            f.write(str(msg).rstrip() + '\n')
    except Exception:
        pass
    logging.error(msg)

def resolve_template(templates_root, clinic, program, filename):
    for p in (
        os.path.join(templates_root, clinic, filename),
        os.path.join(templates_root, 'default', program, filename),
        os.path.join(templates_root, 'default', 'BIPP', filename),
    ):
        if os.path.isfile(p):
            return p
    return None


# ------------------------------
# 5) Template Path via --templates_dir
# ------------------------------
template_path = resolve_template(_TEMPLATES_ROOT, clinic_folder, program_code, 'Template.AC Entrance Notification.docx')
if not template_path:
    record_error(f"[FATAL] Missing template Template.AC Entrance Notification.docx under clinic/default fallbacks in {_TEMPLATES_ROOT}.")


def format_date(value, formats=('%m/%d/%Y','%Y-%m-%d')):
    """Try multiple date formats; return a standard mm/dd/YYYY if successful."""
    if not value or not isinstance(value, str):
        return value
    value = value.strip()
    for fmt in formats:
        try:
            dt = datetime.strptime(value, fmt)
            return dt.strftime('%m/%d/%Y')
        except ValueError:
            continue
    return value

def fill_placeholders(row):
    """
    Clean up 'nan' fields and set placeholders:
      - If any field is 'nan', replace with empty ''
      - Set gender placeholders
      - Format any *_date fields
      - Format orientation_date
    """
    # 1. Replace 'nan' with empty strings
    for col, val in row.items():
        if isinstance(val, str) and val.strip().lower() == 'nan':
            row[col] = ''

    # 2. Gender placeholders
    gender_mapping = {
        'male':   {'gender1': 'his', 'gender2': 'he', 'gender3': "Men's"},
        'female': {'gender1': 'her', 'gender2': 'she', 'gender3': "Women's"}
    }
    gender = row.get('gender', '').lower()
    if gender in gender_mapping:
        row.update(gender_mapping[gender])
    else:
        logging.error(f"Gender '{gender}' not recognized. Setting placeholders to empty.")
        row['gender1'] = ''
        row['gender2'] = ''
        row['gender3'] = ''

    # 3. Convert fields ending with "_date"
    for col in row.keys():
        if col.endswith('_date') and row[col].strip():
            row[col] = format_date(row[col])

    # 4. Force 'dob' → mm/dd/yyyy
    if 'dob' in row and row['dob'].strip():
        row['dob'] = format_date(row['dob'])

    # 5. Force 'orientation_date' → mm/dd/yyyy
    if 'orientation_date' in row and row['orientation_date'].strip():
        row['orientation_date'] = format_date(row['orientation_date'])

    return row

def sanitize_filename(filename):
    invalid_chars = '<>:"/\\|?*'
    for ch in invalid_chars:
        filename = filename.replace(ch, '')
    return filename.strip()

def docx_to_pdf(docx_path, pdf_dir, retries=3):
    """
    Convert DOCX to PDF using LibreOffice with retry logic.
    """
    attempt = 0
    while attempt < retries:
        try:
            output_file = os.path.join(pdf_dir, os.path.basename(docx_path).replace('.docx', '.pdf'))
            subprocess.run(
                ['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', pdf_dir, docx_path],
                check=True
            )
            if os.path.exists(output_file):
                logging.info(f"Converted {docx_path} to PDF successfully on attempt {attempt + 1}.")
                return output_file
            else:
                raise FileNotFoundError(f"PDF not found at expected path: {output_file}")
        except subprocess.CalledProcessError as e:
            logging.error(f"Error converting {docx_path} to PDF on attempt {attempt + 1}: {e}")
            attempt += 1
            time.sleep(2)  # wait 2 seconds before retry
    logging.error(f"Failed to convert {docx_path} to PDF after {retries} attempts.")
    return None

# 1) Load CSV
try:
    df = pd.read_csv(csv_file_path)
    logging.info(f"Loaded {len(df)} records from CSV: {csv_file_path}")
except Exception as e:
    logging.error(f"Failed to load CSV: {e}")
    sys.exit(1)

# 2) Filter: Anger Control + Not Exited
df = df[(df['program_name'] == 'Anger Control') & (df['exit_reason'] == 'Not Exited')]

# 3) Parse orientation_date as datetime
df['orientation_date_str'] = df['orientation_date']  # Keep original string for placeholders
df['orientation_date'] = pd.to_datetime(df['orientation_date'], errors='coerce')

# 4) Filter by orientation_date between start_date & end_date
df = df[(df['orientation_date'] >= start_date) & (df['orientation_date'] <= end_date)]

if df.empty:
    logging.warning("No records found after filtering. No documents will be generated.")
    sys.exit(0)

# Convert to str, fill placeholders
df.drop(columns=['orientation_date'], inplace=True)
df = df.astype(str)
df = df.apply(fill_placeholders, axis=1)
df.fillna('', inplace=True)

def generate_documents_for_row(row):
    if not template_path:
        return  # soft-skip if there’s no template

    """
    Render the AC Entrance Notification for one row, convert to PDF,
    and move it to the correct folder (handling 'nan' for office/manager name).
    """
    try:
        doc = DocxTemplate(template_path)
        context = row.to_dict()

        # orientation_date from orientation_date_str
        orientation_str = context.get('orientation_date_str', '').strip()
        context['orientation_date'] = format_date(orientation_str)
        context.pop('orientation_date_str', None)

        # If needed, re-assign gender placeholders
        gender = context.get('gender', '').lower()
        if gender == 'male':
            context['gender1'] = 'his'
            context['gender2'] = 'he'
            context['gender3'] = "Men's"
        elif gender == 'female':
            context['gender1'] = 'her'
            context['gender2'] = 'she'
            context['gender3'] = "Women's"
        else:
            logging.error(f"Gender '{gender}' not recognized. Setting placeholders empty.")
            context['gender1'] = ''
            context['gender2'] = ''
            context['gender3'] = ''

        doc.render(context)

        first_name = sanitize_filename(context.get('first_name', ''))
        last_name  = sanitize_filename(context.get('last_name', ''))
        docx_filename = f"{last_name} {first_name} AC Entrance.docx".strip()
        docx_path = os.path.join(generated_documents_dir, docx_filename)

        doc.save(docx_path)
        logging.info(f"Document saved at: {docx_path}")

        # Convert to PDF
        pdf_path = docx_to_pdf(docx_path, generated_documents_dir)
        if pdf_path:
            # Remove DOCX after conversion
            os.remove(docx_path)

            # -------------- Office / Case Manager Fields --------------
            case_manager_office     = str(context.get('case_manager_office', '')).strip()
            case_manager_first_name = str(context.get('case_manager_first_name', '')).strip()
            case_manager_last_name  = str(context.get('case_manager_last_name', '')).strip()

            # If office is missing/'nan', place PDF in central_location_path
            if not case_manager_office or case_manager_office.lower() == 'nan':
                office_dir = central_location_path
            else:
                office_dir = os.path.join(
                    central_location_path,
                    sanitize_filename(case_manager_office)
                )

            # If first name is missing/'nan', use only last name for subfolder
            if not case_manager_first_name or case_manager_first_name.lower() == 'nan':
                manager_subfolder_name = sanitize_filename(case_manager_last_name)
            else:
                manager_subfolder_name = f"{sanitize_filename(case_manager_first_name)} {sanitize_filename(case_manager_last_name)}"

            manager_subfolder_name = manager_subfolder_name.strip()
            manager_dir = os.path.join(office_dir, manager_subfolder_name)
            os.makedirs(manager_dir, exist_ok=True)

            # Move PDF
            final_pdf_path = os.path.join(manager_dir, os.path.basename(pdf_path))
            shutil.move(pdf_path, final_pdf_path)
            logging.info(f"Moved {pdf_path} to {final_pdf_path}")
        else:
            logging.error("PDF conversion failed, document not moved.")
    except Exception as e:
        logging.error(f"Error processing document for row: {e}")

# Generate documents for each row
try:
    for _, row in df.iterrows():
        generate_documents_for_row(row)
except Exception as e:
    logging.error(f"Error during document generation loop: {e}")
    sys.exit(1)

logging.info("AC Entrance Notifications have been generated and organized.")
print("AC Entrance Notifications have been generated and organized.")
