#!/usr/bin/env python3
import pandas as pd
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Inches
import logging
import requests
from io import BytesIO
import os
import sys
import mysql.connector
from datetime import datetime, timedelta
from PIL import Image
import argparse
import shutil
import subprocess
from docx.enum.table import WD_TABLE_ALIGNMENT
import time

################################################################################
# ARGUMENT PARSING
################################################################################

parser = argparse.ArgumentParser(description='Generate BIPP Curriculum Progress Reports')

parser.add_argument('--csv_file', required=True, help='Path to the CSV file')
parser.add_argument('--start_date', required=False, help='(Ignored) Start date')
parser.add_argument('--end_date', required=False, help='(Ignored) End date')
parser.add_argument('--output_dir', required=False, help='Path to the output directory')
parser.add_argument('--clinic_folder', required=True, help='Clinic folder identifier')
parser.add_argument('--templates_dir', required=True, help='Path to the templates directory')
# DB credentials
parser.add_argument('--db_host', required=True, help='Database hostname')
parser.add_argument('--db_user', required=True, help='Database username')
parser.add_argument('--db_pass', required=True, help='Database password')
parser.add_argument('--db_name', required=True, help='Database name')

args = parser.parse_args()

csv_file_path   = args.csv_file
output_dir      = args.output_dir
clinic_folder   = args.clinic_folder
templates_dir   = args.templates_dir
db_host         = args.db_host
db_user         = args.db_user
db_pass         = args.db_pass
db_name         = args.db_name

# --- Env overrides to match other BIPP scripts ---
program_code  = os.environ.get('NP_PROGRAM_CODE', 'BIPP')
templates_dir = os.environ.get('NP_TEMPLATES_DIR', templates_dir)
clinic_folder = os.environ.get('NP_CLINIC_FOLDER', clinic_folder)

# after env overrides
if os.path.basename(templates_dir) == clinic_folder:
    _TEMPLATES_ROOT = os.path.dirname(templates_dir)
else:
    _TEMPLATES_ROOT = templates_dir

################################################################################
# LOGGING SETUP
################################################################################

log_dir = '/home/notesao/NotePro-Report-Generator/logs'
os.makedirs(log_dir, exist_ok=True)
logging.basicConfig(
    filename=os.path.join(log_dir, f'BIPP_CUR_Progress_Reports_{datetime.now().strftime("%Y%m%d")}.log'),
    level=logging.DEBUG,
    format='%(asctime)s [%(levelname)s] %(message)s'
)
logging.info("BIPP Curriculum Progress Reports Script started.")

################################################################################
# FOLDER SETUP
################################################################################

base_dir = os.path.dirname(os.path.abspath(__file__))

today_str = datetime.now().strftime('%m.%d.%y')
if output_dir:
    generated_documents_dir = output_dir
else:
    generated_documents_dir = os.path.join(base_dir, 'GeneratedDocuments', clinic_folder, today_str)
os.makedirs(generated_documents_dir, exist_ok=True)

public_generated_documents_dir = f"/home/notesao/{clinic_folder}/public_html/GeneratedDocuments"
os.makedirs(public_generated_documents_dir, exist_ok=True)

central_location_path = os.path.join(public_generated_documents_dir, today_str)
os.makedirs(central_location_path, exist_ok=True)
logging.info(f"Central location directory created: {central_location_path}")

# Per-run human-friendly error sink
errors_txt = os.path.join(central_location_path, 'errors.txt')

def record_error(msg: str):
    try:
        with open(errors_txt, 'a', encoding='utf-8') as f:
            f.write(msg.rstrip() + '\n')
    except Exception as e:
        logging.error(f"Failed writing to errors.txt: {e}")
    logging.warning(msg)

def ensure_columns(df, cols, default=''):
    for c in cols:
        if c not in df.columns:
            df[c] = default
    return df


################################################################################
# TEMPLATES
################################################################################

def resolve_template(templates_root: str, clinic: str, program: str, filename: str):
    candidates = [
        os.path.join(templates_root, clinic, filename),
        os.path.join(templates_root, 'default', program, filename),
        os.path.join(templates_root, 'default', 'BIPP', filename),
    ]
    for p in candidates:
        if os.path.isfile(p):
            return p
    return None

def need_template(filename):
    path = resolve_template(_TEMPLATES_ROOT, clinic_folder, program_code, filename)
    if not path:
        record_error(f"[FATAL] Missing template {filename} under clinic/default fallbacks in {_TEMPLATES_ROOT}.")
        sys.exit(0)
    return path

def try_template(filename):
    p = resolve_template(_TEMPLATES_ROOT, clinic_folder, program_code, filename)
    if not p:
        record_error(f"[WARN] Missing optional template {filename}; falling back to standard template.")
    return p

progress_report_template_path         = need_template('Template.BIPP CUR Progress Report.docx')
virtual_progress_report_template_path = try_template('Template.BIPP CUR Progress Report Virtual.docx')


################################################################################
# DB CONNECTION
################################################################################

try:
    db_conn = mysql.connector.connect(
        host=db_host,
        user=db_user,
        password=db_pass,
        database=db_name
    )
    db_cursor = db_conn.cursor(dictionary=True)
    logging.info("Database connection established successfully.")
except Exception as e:
    logging.error(f"Failed to connect to DB: {e}")
    sys.exit(1)

################################################################################
# LOAD CURRICULUM BLANK NOTES (Excel)
################################################################################

def find_curriculum_excel(root, clinic, program):
    for p in (
        os.path.join(root, clinic, "BIPP Curriculum Blank Notes.xlsx"),
        os.path.join(root, 'default', program, "BIPP Curriculum Blank Notes.xlsx"),
        os.path.join(root, 'default', 'BIPP', "BIPP Curriculum Blank Notes.xlsx"),
    ):
        if os.path.isfile(p):
            return p
    return None

curriculum_notes_path = find_curriculum_excel(_TEMPLATES_ROOT, clinic_folder, program_code)

# Ensure we always have a dataframe, even on failure
curriculum_df = pd.DataFrame()



try:
    if not curriculum_notes_path:
        raise FileNotFoundError("BIPP Curriculum Blank Notes.xlsx not found under clinic/default fallbacks")
    curriculum_df = pd.read_excel(curriculum_notes_path)
    curriculum_df.columns = [c.strip() for c in curriculum_df.columns]
    logging.info(f"Loaded curriculum notes: {len(curriculum_df)} rows from {curriculum_notes_path}")
except Exception as e:
    logging.warning(f"Curriculum Excel load failed ({e}); continuing without curriculum notes.")
    curriculum_df = pd.DataFrame(columns=[
        "Curriculum Name", "Theme", "Theme Name", "Part", "Part Name", "Part Note"
    ])


# --- build normalized title index (MUST be after loading curriculum_df) ---
NAME_CANDIDATES = ['Part Name','Curriculum Name','Topic','Module','Name']
INDEX_COLS = [c for c in NAME_CANDIDATES if c in curriculum_df.columns]
if not INDEX_COLS:
    record_error(f"[FATAL] No curriculum title column found. Tried: {NAME_CANDIDATES}")
    sys.exit(0)

import unicodedata, re
def _norm(s):
    s = '' if pd.isna(s) else str(s)
    s = unicodedata.normalize('NFKC', s)
    s = s.replace('\u2018',"'").replace('\u2019',"'")
    s = s.replace('\u2013','-').replace('\u2014','-')
    s = re.sub(r'\s+',' ', s).strip().lower()
    return s

# Prefer matches from Part Name over Curriculum Name, then others
PREF_RANK = {col: rank for rank, col in enumerate(INDEX_COLS)}  # lower rank = more preferred
TITLE_INDEX = {}  # norm_title -> (rank, row_index, matched_col)

for i, r in curriculum_df.iterrows():
    for col in INDEX_COLS:
        val = str(r.get(col, '')).strip()
        if not val:
            continue
        k = _norm(val)
        if k not in TITLE_INDEX or PREF_RANK[col] < TITLE_INDEX[k][0]:
            TITLE_INDEX[k] = (PREF_RANK[col], i, col)

from difflib import get_close_matches
def find_curriculum_row_by_title(title: str):
    t = _norm(title)
    if t in TITLE_INDEX:
        _, idx, col = TITLE_INDEX[t]
        logging.debug(f"Mapped '{title}' via exact match on column '{col}' (row {idx}).")
        return curriculum_df.iloc[idx]
    # fuzzy over all keys we indexed
    keys = list(TITLE_INDEX.keys())
    matches = get_close_matches(t, keys, n=1, cutoff=0.78)
    if matches:
        _, idx, col = TITLE_INDEX[matches[0]]
        logging.debug(f"Mapped '{title}' via fuzzy match to '{matches[0]}' on column '{col}' (row {idx}).")
        return curriculum_df.iloc[idx]
    # last-chance: substring contains across all index columns
    for i, r in curriculum_df.iterrows():
        for col in INDEX_COLS:
            if t and t in _norm(str(r.get(col, ''))):
                logging.debug(f"Mapped '{title}' via substring match on column '{col}' (row {i}).")
                return curriculum_df.iloc[i]
    return None




################################################################################
# HELPER FUNCTIONS
################################################################################

def docx_to_pdf(docx_path, pdf_dir, retries=3):
    attempt = 0
    while attempt < retries:
        try:
            output_file = os.path.join(pdf_dir, os.path.basename(docx_path).replace('.docx', '.pdf'))
            subprocess.run(
                ['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', pdf_dir, docx_path],
                check=True
            )
            if os.path.exists(output_file):
                logging.info(f"Converted {docx_path} to PDF successfully on attempt {attempt + 1}.")
                return output_file
            else:
                raise FileNotFoundError(f"PDF not found at {output_file}")
        except FileNotFoundError as e:
            record_error("LibreOffice not found on PATH. Install it or add it to PATH to enable DOCX→PDF.")
            break
        except subprocess.CalledProcessError as e:
            logging.error(f"Error converting {docx_path} to PDF on attempt {attempt + 1}: {e}")
            attempt += 1
            time.sleep(2)
    logging.error(f"Failed to convert {docx_path} to PDF after {retries} attempts.")
    return None


def align_tables_left(doc):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.LEFT

def fetch_image(url, doc, max_width=3.45, max_height=4.71, margin=0.2):
    try:
        if not isinstance(url, str) or not url.lower().startswith(('http://','https://')):
            logging.warning(f"Skipping non-http image_url: {url}")
            return None

        headers = {
            'User-Agent': 'Mozilla/5.0',
            'Accept': 'image/*'
        }
        response = requests.get(url, headers=headers, timeout=30)
        if response.status_code == 200 and 'image' in response.headers.get('Content-Type', ''):
            img = Image.open(BytesIO(response.content))

            img_width, img_height = img.size
            aspect_ratio = img_width / img_height

            adjusted_max_width  = max_width - margin
            adjusted_max_height = max_height - margin

            # Resize to fit
            if aspect_ratio > 1:
                new_width  = min(adjusted_max_width, img_width)
                new_height = new_width / aspect_ratio
                if new_height > adjusted_max_height:
                    new_height = adjusted_max_height
                    new_width  = new_height * aspect_ratio
            else:
                new_height = min(adjusted_max_height, img_height)
                new_width  = new_height * aspect_ratio
                if new_width > adjusted_max_width:
                    new_width  = adjusted_max_width
                    new_height = new_width / aspect_ratio

            img = img.resize((int(new_width * 300), int(new_height * 300)), Image.LANCZOS)
            img_byte_arr = BytesIO()
            img.save(img_byte_arr, format='JPEG')
            return InlineImage(doc, BytesIO(img_byte_arr.getvalue()), width=Inches(new_width), height=Inches(new_height))
        else:
            logging.warning(f"Failed to fetch image or non-image content returned: {url}")
            return None
    except Exception as e:
        logging.error(f"Unexpected error fetching image: {e}")
        return None

def convert_yn_to_yes_no(value):
    if isinstance(value, str):
        val = value.strip().lower()
        if val == 'y':
            return 'Yes'
        elif val == 'n':
            return 'No'
    return value

def parse_and_format_date(date_str):
    """ Format date_str as m/d/yyyy if parseable. """
    if not date_str or not isinstance(date_str, str):
        return ''
    date_str = date_str.strip()
    if not date_str:
        return ''
    for fmt in ('%Y-%m-%d', '%m/%d/%Y'):
        try:
            d = datetime.strptime(date_str, fmt)
            return f"{d.month}/{d.day}/{d.year}"  # m/d/yyyy
        except ValueError:
            continue
    return date_str

def parse_and_format_p_a_date(date_str):
    """ Format P# / A# date as mm/dd/yy if parseable. """
    if not date_str or not isinstance(date_str, str):
        return ''
    date_str = date_str.strip()
    if not date_str:
        return ''
    for fmt in ('%Y-%m-%d', '%m/%d/%Y'):
        try:
            d = datetime.strptime(date_str, fmt)
            return d.strftime('%m/%d/%y')  # mm/dd/yy
        except ValueError:
            continue
    return ''

def determine_fee_problems(balance):
    """ Return 'No' if balance==0 else 'Yes' """
    try:
        b = float(balance)
        if b == 0:
            return "No"
        return "Yes"
    except:
        return ""

def format_balance(balance):
    """ Format as e.g. $40. """
    try:
        b = float(balance)
        return f"${int(b)}"
    except:
        return "$0"

def determine_stage_of_change(attended_str, required_str):
    """Same logic as your Stage-of-Change script."""
    try:
        attended = float(attended_str)
        required = float(required_str)
    except ValueError:
        logging.error(f"Invalid attended/required: {attended_str}, {required_str}")
        return 'Alert'

    if required == 0:
        return 'Alert'

    # 18-week logic
    if required == 18:
        if 1 <= attended <= 4:
            return 'Precontemplation'
        elif 5 <= attended <= 8:
            return 'Contemplation'
        elif 9 <= attended <= 11:
            return 'Preparation'
        elif 12 <= attended <= 15:
            return 'Action'
        elif 16 <= attended <= 18:
            return 'Maintenance'
        else:
            logging.warning(f"Unexpected attendance for 18-week: {attended}")
            return 'Alert'

    # 27-week logic
    if required == 27:
        if 1 <= attended <= 7:
            return 'Precontemplation'
        elif 8 <= attended <= 14:
            return 'Contemplation'
        elif 15 <= attended <= 18:
            return 'Preparation'
        elif 19 <= attended <= 23:
            return 'Action'
        elif 24 <= attended <= 27:
            return 'Maintenance'
        else:
            logging.warning(f"Unexpected attendance for 27-week: {attended}")
            return 'Alert'

    # Adaptive logic for other durations
    ratio = attended / required
    if ratio > 0   and ratio <= (4/18):
        return 'Precontemplation'
    elif ratio <= (8/18):
        return 'Contemplation'
    elif ratio <= (11/18):
        return 'Preparation'
    elif ratio <= (15/18):
        return 'Action'
    elif ratio <= 1.0:
        return 'Maintenance'
    else:
        logging.warning(f"Ratio out of expected range for required={required}, attended={attended}")
        return 'Alert'

def fill_placeholders(row):
    """
    We replicate the SOC logic so that DOB, orientation_date, last_attended,
    P/A fields, D# checkmarks, stage_of_change, fee, etc. all work the same.
    """
    # Gender placeholders
    gender = row.get('gender', '').lower()
    if gender == 'male':
        row['gender1'] = 'his'
        row['gender2'] = 'he'
        row['gender3'] = "Men's"
        row['gender4'] = 'him'
        row['gender5'] = "He"
        row['gender6'] = 'himself'
    elif gender == 'female':
        row['gender1'] = 'her'
        row['gender2'] = 'she'
        row['gender3'] = "Women's"
        row['gender4'] = 'her'
        row['gender5'] = "She"
        row['gender6'] = 'herself'
    else:
        # neutral fallbacks
        row['gender1'] = 'their'
        row['gender2'] = 'they'
        row['gender3'] = "Participants"
        row['gender4'] = 'them'
        row['gender5'] = "They"
        row['gender6'] = 'themself'


    # Format key dates
    for date_field in ['dob', 'orientation_date', 'last_attended']:
        if date_field in row:
            row[date_field] = parse_and_format_date(row[date_field])

    # Mark D1..D30
    attended = row.get('attended', '0')
    required = row.get('required_sessions', '0')
    try:
        attended_f = float(attended)
        required_f = float(required)
    except ValueError:
        attended_f, required_f = 0, 0

    for i in range(1, 31):
        row[f'D{i}'] = ''

    # D1 is always a check
    row['D1'] = '\u2714'

    if required_f > 0:
        filled = int((attended_f / required_f) * 30)
        for i in range(2, filled + 1):
            if i <= 30:
                row[f'D{i}'] = '\u2714'

    # If there's a CSV-based client_note, replace & with 'and'
    if 'client_note' in row and isinstance(row['client_note'], str):
        row['client_note'] = row['client_note'].replace('&', 'and')

    # P1..P27, A1..A18 => mm/dd/yy
    for i in range(1, 28):
        p_field = f'P{i}'
        if p_field in row:
            row[p_field] = parse_and_format_p_a_date(row[p_field])

    for i in range(1, 19):
        a_field = f'A{i}'
        if a_field in row:
            row[a_field] = parse_and_format_p_a_date(row[a_field])

    # Stage of change
    row['client_stagechange'] = determine_stage_of_change(attended, required)

    return row

def sanitize_filename(filename):
    invalid_chars = '<>:"/\\|?*'
    for char in invalid_chars:
        filename = filename.replace(char, '')
    return filename.strip()

################################################################################
# CSV LOADING + CLEANUP
################################################################################

try:
    df = pd.read_csv(csv_file_path, dtype=str)
    # normalize column names (trim stray spaces)
    df.columns = df.columns.str.strip()
    logging.info(f"Loaded {len(df)} records from CSV.")

except Exception as e:
    logging.error(f"Failed to load CSV file: {e}")
    sys.exit(1)

BASE_REQUIRED = [
    'client_id','first_name','last_name','program_name','exit_reason',
    'orientation_date','last_attended','attended','required_sessions'
]

OPTIONAL_CUR_ORG_COLS = [
    'facilitator_office','facilitator_first_name','facilitator_last_name',
    'instructor_office','instructor_first_name','instructor_last_name'
]

if clinic_folder.lower() in ('dwag', 'ffltest'):
    REQUIRED_COLUMNS = BASE_REQUIRED               # org columns optional
else:
    REQUIRED_COLUMNS = BASE_REQUIRED + OPTIONAL_CUR_ORG_COLS

missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
if missing:
    if clinic_folder.lower() in ('dwag', 'ffltest'):
        # fill any missing as blanks and continue
        df = ensure_columns(df, missing, '')
        logging.info(f"Optional org columns missing for {clinic_folder}; continuing with blanks: {missing}")
    else:
        record_error(f"Missing required columns in CSV: {missing}. Exiting non-fatally.")
        sys.exit(0)

df.fillna('', inplace=True)  # everything is string, so blank is fine

# DWAG/ffltest do not require these — create blank columns if missing
if clinic_folder.lower() in ('dwag', 'ffltest'):
    df = ensure_columns(df, OPTIONAL_CUR_ORG_COLS, '')


# Convert Y/N to Yes/No only on object (string-like) columns
for col in df.select_dtypes(include='object').columns:
    df[col] = df[col].apply(convert_yn_to_yes_no)


# Normalize and filter (robust to casing/whitespace)
df['program_name']   = df['program_name'].astype(str).str.strip()
df['exit_reason']    = df['exit_reason'].astype(str).str.strip()
df['exit_reason_lc'] = df['exit_reason'].str.lower()

df = df[df['program_name'].isin(['BIPP (male)', 'BIPP (female)'])]
df = df[df['exit_reason_lc'] == 'not exited']

if df.empty:
    logging.warning("No BIPP Not Exited clients found. Exiting.")
    sys.exit(0)

df = ensure_columns(df, ['exit_date', 'balance'], '')
# Exit date handling (datetime → string yyyy-mm-dd)
df['exit_date'] = pd.to_datetime(df['exit_date'], errors='coerce')
df['exit_date'] = df['exit_date'].dt.strftime('%Y-%m-%d').fillna('')

# If last_attended is blank, fall back to orientation_date
def fix_last_attended(row):
    la = str(row.get('last_attended', '')).strip()
    return row.get('orientation_date', '') if not la else la

df['last_attended'] = df.apply(fix_last_attended, axis=1)

# Fee/Balance fields
df['feeproblems'] = df['balance'].apply(determine_fee_problems)
df['balance']     = df['balance'].apply(format_balance)

# Coerce numeric fields used for stage-of-change/checkbox logic
df['attended']          = pd.to_numeric(df['attended'], errors='coerce').fillna(0)
df['required_sessions'] = pd.to_numeric(df['required_sessions'], errors='coerce').fillna(0)

# Apply placeholders (DOB, orientation_date, D# checks, P/A dates, etc.)
df = df.apply(fill_placeholders, axis=1)

# --- DWAG-only derived placeholders (no DB changes) -------------------
def _as_bool(v):
    s = str(v).strip().lower()
    return s in ('yes','y','true','1','✓','check','checked')

def _chk(ok):  # ✓ / ✗
    return '\u2714' if ok else '\u2716'

def _dwag_derived(row: pd.Series) -> pd.Series:
    # Absence count
    try:
        absences = int(float(str(row.get('absence_unexcused','0')).strip() or 0))
    except Exception:
        absences = 0

    # Attended initial assessment (orientation_date present)
    has_orientation = bool(str(row.get('orientation_date','')).strip())
    row['attended_initial_assessment']  = _chk(has_orientation)
    row['attended_initial_assessessment'] = row['attended_initial_assessment']  # template typo compatibility

    # Consistent attendance (no absences)
    row['consistent_attendance'] = _chk(absences == 0)

    # Three or more absences text
    if absences >= 3:
        row['three_or_more_absences'] = f"{absences} absences"
    elif absences == 0:
        row['three_or_more_absences'] = "No"
    else:
        row['three_or_more_absences'] = "Within 3"

    # Participation → compliance checks
    respectful = _as_bool(row.get('respectful_to_group','No'))
    disruptive = _as_bool(row.get('disruptive_arumentitive','No')) or _as_bool(row.get('disruptive_argumentative','No'))  # handle typo
    humor_bad  = _as_bool(row.get('humor_inappropriate','No'))
    blames     = _as_bool(row.get('blames_victim','No'))
    appears    = _as_bool(row.get('appears_drug_alcohol','No'))
    staff_bad  = _as_bool(row.get('inappropriate_to_staff','No'))
    dialogue   = _as_bool(row.get('participated_in_dialogue','')) or _as_bool(row.get('participated_in_dialoge',''))
    takes_resp = _as_bool(row.get('takes_responsibility_for_past','')) or _as_bool(row.get('takes_responsibility_for past',''))

    row['followed_group_rules']            = _chk(respectful and not disruptive and not staff_bad and not humor_bad)
    row['adhered_to_participant_agreement']= _chk(not appears and not blames and not staff_bad)
    row['participated_in_dialogue']        = _chk(dialogue or (respectful and not disruptive))
    row['participated_in_dialoge']         = row['participated_in_dialogue']  # template typo compatibility
    row['accountability_consistent']       = _chk(takes_resp)
    return row

if clinic_folder.lower() == 'dwag':
    df = df.apply(_dwag_derived, axis=1)
# ----------------------------------------------------------------------


################################################################################
# DB + EXCEL => CLIENT_NOTE
################################################################################

def fetch_curriculum_note_for_client(client_id):
    """
    Fetches the curriculum notes based on the earliest curriculum taught in the last 22 days.
    If a single curriculum is found, it infers the next two in sequence.
    """
    if curriculum_df.empty:
        return ""

    today = datetime.now().date()
    start_date = today - timedelta(days=22)
    end_date   = today  # exclusive

    sql = """
        SELECT ts.date AS session_date, ts.curriculum_id AS cid
          FROM attendance_record ar
          JOIN therapy_session ts ON ar.therapy_session_id = ts.id
         WHERE ar.client_id = %s
           AND ts.date >= %s
           AND ts.date < %s
           AND ts.curriculum_id IS NOT NULL
           AND ts.curriculum_id <> ''
         ORDER BY ts.date ASC
    """
    db_cursor.execute(sql, (client_id, start_date, end_date))
    sessions = db_cursor.fetchall()

    if not sessions:
        return ""

    # Earliest session => earliest_cid => earliest long_description
    earliest_cid = sessions[0]['cid']
    db_cursor.execute("SELECT long_description FROM curriculum WHERE id = %s", (earliest_cid,))
    row_first = db_cursor.fetchone()
    if not row_first or not row_first['long_description']:
        return ""
    
    first_sd = row_first['long_description'].strip()

    # Identify curriculum sequence
    base_row = find_curriculum_row_by_title(first_sd)   # db_title_string == first_sd
    if base_row is None:
        logging.warning(f"Could not map '{first_sd}' to a curriculum row.")
        return ""

    base_idx  = base_row.name
    num_total = len(curriculum_df)
    wanted_idxs = [(base_idx + offset) % num_total for offset in range(3)]

    final_lines = []
    for i, ridx in enumerate(wanted_idxs, start=1):
        subrow    = curriculum_df.iloc[ridx]
        part_name = str(subrow.get('Part Name', '')).strip()
        part_note = str(subrow.get('Part Note', '')).strip()
        if part_name:
            final_lines.append(f"{part_name} (Week {i})")
        if part_note:
            final_lines.append(part_note)
        final_lines.append("")


    return "\n".join(final_lines).strip()



def render_and_save(template_path, doc_type, context, output_dir):
    try:
        doc = DocxTemplate(template_path)
        # Insert image if available
        if 'image_url' in context and context['image_url']:
            image = fetch_image(context['image_url'], doc)
            context['image'] = image if image else ''

        doc.render(context)
        align_tables_left(doc)

        first_name = sanitize_filename(context.get('first_name', ''))
        last_name  = sanitize_filename(context.get('last_name',  ''))
        docx_filename = f"{last_name} {first_name} BIPP Curriculum PReport.docx"
        docx_path = os.path.join(output_dir, docx_filename)

        doc.save(docx_path)
        time.sleep(1)

        pdf_path = docx_to_pdf(docx_path, output_dir)
        if pdf_path:
            os.remove(docx_path)
            return pdf_path
        else:
            logging.error(f"PDF conversion failed for {docx_path}")
            return None
    except Exception as e:
        logging.error(f"Error rendering {doc_type}: {e}")
        record_error(f"Render failed for {doc_type} / {context.get('last_name','')}, {context.get('first_name','')}: {e}")
        return None


def generate_documents(row, output_dir_path):
    # We already have row placeholders from fill_placeholders
    context = row.to_dict()

    # ------------------------------------------------------------------
    # 1) Compute friendly date range for {{client_note_dates}} 
    #    e.g. "January 14th, 2025 through February 4th, 2025"
    # ------------------------------------------------------------------
    # Usually you want 22 days ago as the start and "yesterday" as the end
    start_dt = datetime.now().date() - timedelta(days=22)  # e.g. Jan 14
    end_dt   = datetime.now().date() - timedelta(days=1)   # e.g. Feb 4
    
    def format_friendly_date(dt):
        """Optional: add day-number suffix (1st, 2nd, 3rd, etc.)"""
        # Quick suffix logic:
        day = dt.day
        if 11 <= day % 100 <= 13:
            suffix = "th"
        elif day % 10 == 1:
            suffix = "st"
        elif day % 10 == 2:
            suffix = "nd"
        elif day % 10 == 3:
            suffix = "rd"
        else:
            suffix = "th"
        return dt.strftime(f"%B {day}{suffix}, %Y")

    start_str = format_friendly_date(start_dt)
    end_str   = format_friendly_date(end_dt)
    context["client_note_dates"] = f"{start_str} through {end_str}"

    # Overwrite 'client_note' from DB + Excel
    client_id_str = str(context.get("client_id", "")).strip()

    # normalize common float-ish IDs like "123.0"
    if client_id_str.endswith(".0"):
        client_id_str = client_id_str[:-2]

    # as a last resort, keep only digits (e.g., " ID:123 ")
    if not client_id_str.isdigit():
        digits_only = ''.join(ch for ch in client_id_str if ch.isdigit())
        client_id_str = digits_only

    if not client_id_str.isdigit():
        logging.warning(f"Row has invalid client_id={context.get('client_id', '')}. Skipping DB logic.")
        context["client_note"] = ""
    else:
        client_id = int(client_id_str)
        # Fetch final note from the function below
        final_note = fetch_curriculum_note_for_client(client_id)


        # ---------------------------------------------------------
        # Manual string replacement for the placeholders 
        # ({{first_name}}, {{gender1}}, etc.)
        # ---------------------------------------------------------
        replacements = {
            "{{first_name}}": context.get("first_name", ""),
            "{{gender1}}":    context.get("gender1", ""),
            "{{gender2}}":    context.get("gender2", ""),
            "{{gender3}}":    context.get("gender3", ""),
            "{{gender4}}":    context.get("gender4", ""),
            "{{gender5}}":    context.get("gender5", ""),
            "{{gender6}}":    context.get("gender6", ""),
        }
        for placeholder, value in replacements.items():
            final_note = final_note.replace(placeholder, value)

        context["client_note"] = final_note

    # The doc uses "report_date" for today's date
    now = datetime.now()
    context['report_date'] = f"{now.month}/{now.day}/{now.year}"

    # Pick the appropriate Word template
    group_name = str(context.get("group_name", "")).lower()
    report_template = progress_report_template_path
    if (("virtual" in group_name) or ("women" in group_name)) and virtual_progress_report_template_path:
        report_template = virtual_progress_report_template_path



    pdf_path = render_and_save(report_template, "PReport", context, output_dir_path)

    # Move PDF into case manager folder
    # Move PDF into case manager / facilitator / instructor folder (clinic differences)
    def pick(d, *keys, default=''):
        for k in keys:
            v = d.get(k, '')
            if isinstance(v, str) and v.strip():
                return v.strip()
        return default

    raw_office = pick(row, 'case_manager_office','facilitator_office','instructor_office', default='')
    office_dir = (output_dir_path if not raw_office or raw_office.lower() in ('nan','none','null','')
                else os.path.join(output_dir_path, sanitize_filename(raw_office)))

    raw_first = pick(row, 'case_manager_first_name','facilitator_first_name','instructor_first_name', default='')
    raw_last  = pick(row, 'case_manager_last_name', 'facilitator_last_name', 'instructor_last_name',  default='')
    manager_subfolder = (f"{sanitize_filename(raw_first)} {sanitize_filename(raw_last)}".strip()
                        or (sanitize_filename(raw_last) if raw_last else 'Unknown Manager'))

    manager_dir = os.path.join(office_dir, manager_subfolder)
    os.makedirs(manager_dir, exist_ok=True)

    if pdf_path and os.path.exists(pdf_path):
        final_path = os.path.join(manager_dir, os.path.basename(pdf_path))
        shutil.move(pdf_path, final_path)
        logging.info(f"Moved {pdf_path} to {final_path}")
        return True
    else:
        logging.warning("PDF file not found or not generated.")
        return False



################################################################################
# MAIN LOOP
################################################################################

processed = 0
skipped   = 0
try:
    for idx, row in df.iterrows():
        try:
            ok = generate_documents(row, central_location_path)
            processed += 1 if ok else 0
            skipped   += 0 if ok else 1
        except Exception as row_e:
            skipped += 1
            record_error(f"[row {idx}] Unexpected error: {row_e}")
            continue
except Exception as e:
    record_error(f"[FATAL] Loop crashed: {e}")
finally:
    try:
        db_cursor.close()
        db_conn.close()
    except Exception:
        pass

logging.info(f"BIPP CUR Progress Reports completed. Processed={processed}, Skipped={skipped}")
print(f"BIPP CUR Progress Reports complete. Processed={processed}, Skipped={skipped}")
sys.exit(0)

