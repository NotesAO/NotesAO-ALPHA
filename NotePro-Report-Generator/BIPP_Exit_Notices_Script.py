import pandas as pd
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Inches
import logging
import requests
from io import BytesIO
import os
import sys
from datetime import datetime
from PIL import Image
import argparse
import shutil
import subprocess
from docx.enum.table import WD_TABLE_ALIGNMENT
import time
import re

parser = argparse.ArgumentParser(description='Generate BIPP Exit Notices')
parser.add_argument('--csv_file', required=True, help='Path to the CSV file')
parser.add_argument('--start_date', required=False, help='Start date for the reports (YYYY-MM-DD)')
parser.add_argument('--end_date', required=False, help='End date for the reports (YYYY-MM-DD)')
parser.add_argument('--output_dir', required=False, help='Path to the output directory')
parser.add_argument('--clinic_folder', required=True, help='Clinic folder identifier')
parser.add_argument('--templates_dir', required=True, help='Path to the templates directory')
args = parser.parse_args()

csv_file_path = args.csv_file
start_date_str = args.start_date
end_date_str = args.end_date
output_dir = args.output_dir
clinic_folder = args.clinic_folder

# --- Env overrides to match the other BIPP scripts ---
program_code  = os.environ.get('NP_PROGRAM_CODE', 'BIPP')
templates_dir = os.environ.get('NP_TEMPLATES_DIR', args.templates_dir)
clinic_folder = os.environ.get('NP_CLINIC_FOLDER', clinic_folder)

# after env overrides
if os.path.basename(templates_dir) == clinic_folder:
    _TEMPLATES_ROOT = os.path.dirname(templates_dir)
else:
    _TEMPLATES_ROOT = templates_dir

log_dir = '/home/notesao/NotePro-Report-Generator/logs'
os.makedirs(log_dir, exist_ok=True)
logging.basicConfig(
    filename=os.path.join(log_dir, f'BIPP_Exit_Notices_{datetime.now().strftime("%Y%m%d")}.log'),
    level=logging.DEBUG,
    format='%(asctime)s [%(levelname)s] %(message)s'
)
logging.info("BIPP Exit Notices Script started.")

base_dir = os.path.dirname(os.path.abspath(__file__))

today = datetime.now().strftime('%m.%d.%y')
if output_dir:
    generated_documents_dir = output_dir
else:
    generated_documents_dir = os.path.join(base_dir, 'GeneratedDocuments', clinic_folder, today)
os.makedirs(generated_documents_dir, exist_ok=True)

public_generated_documents_dir = f"/home/notesao/{clinic_folder}/public_html/GeneratedDocuments"
os.makedirs(public_generated_documents_dir, exist_ok=True)

central_location_path = os.path.join(public_generated_documents_dir, today)
os.makedirs(central_location_path, exist_ok=True)
logging.info(f"Central location directory created: {central_location_path}")

# Per-run human-friendly error sink
errors_txt = os.path.join(central_location_path, 'errors.txt')

def record_error(msg: str):
    try:
        with open(errors_txt, 'a', encoding='utf-8') as f:
            f.write(msg.rstrip() + '\n')
    except Exception as e:
        logging.error(f"Failed writing to errors.txt: {e}")
    logging.warning(msg)


def resolve_template(templates_root: str, clinic: str, program: str, filename: str):
    candidates = [
        os.path.join(templates_root, clinic, filename),
        os.path.join(templates_root, 'default', program, filename),
        os.path.join(templates_root, 'default', 'BIPP', filename),
    ]
    for p in candidates:
        if os.path.isfile(p):
            return p
    return None

def need_template(filename):
    p = resolve_template(_TEMPLATES_ROOT, clinic_folder, program_code, filename)
    if not p:
        record_error(f"[FATAL] Missing template {filename} under clinic/default fallbacks in {_TEMPLATES_ROOT}.")
        sys.exit(0)
    return p

exit_notice_template_path              = need_template('Template.BIPP Exit Notice.docx')
progress_report_template_path          = need_template('Template.BIPP Exit Progress Report.docx')
virtual_progress_report_template_path  = need_template('Template.BIPP Exit Progress Report Virtual.docx')


def docx_to_pdf(docx_path, pdf_dir, retries=3):
    attempt = 0
    while attempt < retries:
        try:
            output_file = os.path.join(pdf_dir, os.path.basename(docx_path).replace('.docx', '.pdf'))
            subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', pdf_dir, docx_path], check=True)
            if os.path.exists(output_file):
                logging.info(f"Converted {docx_path} to PDF successfully on attempt {attempt + 1}.")
                return output_file
            else:
                raise FileNotFoundError(f"PDF not found at expected path: {output_file}")
        except subprocess.CalledProcessError as e:
            logging.error(f"Error converting {docx_path} to PDF on attempt {attempt + 1}: {e}")
            attempt += 1
            time.sleep(2)
    logging.error(f"Failed to convert {docx_path} to PDF after {retries} attempts.")
    return None

def align_tables_left(doc):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.LEFT

def fetch_image(url, doc, max_width=3.45, max_height=4.71, margin=0.2):
    headers = {
        'User-Agent': 'Mozilla/5.0',
        'Accept': 'image/*'
    }
    try:
        if not isinstance(url, str) or not url.lower().startswith(('http://','https://')):
            logging.warning(f"Skipping non-http image_url: {url}")
            return None

        response = requests.get(url, headers=headers, timeout=30)
        if response.status_code == 200 and 'image' in response.headers.get('Content-Type', ''):
            img = Image.open(BytesIO(response.content))

            img_width, img_height = img.size
            aspect_ratio = img_width / img_height

            adjusted_max_width = max_width - margin
            adjusted_max_height = max_height - margin

            if aspect_ratio > 1:
                new_width = min(adjusted_max_width, img_width)
                new_height = new_width / aspect_ratio
                if new_height > adjusted_max_height:
                    new_height = adjusted_max_height
                    new_width = new_height * aspect_ratio
            else:
                new_height = min(adjusted_max_height, img_height)
                new_width = new_height * aspect_ratio
                if new_width > adjusted_max_width:
                    new_width = adjusted_max_width
                    new_height = new_width / aspect_ratio

            img = img.resize((int(new_width * 300), int(new_height * 300)), Image.LANCZOS)

            img_byte_arr = BytesIO()
            img.save(img_byte_arr, format='JPEG')

            return InlineImage(doc, BytesIO(img_byte_arr.getvalue()), width=Inches(new_width), height=Inches(new_height))
        else:
            logging.warning("Failed to fetch image or non-image content returned: %s", url)
            return None
    except requests.RequestException as e:
        logging.warning("Request failed for URL %s with error: %s", url, e)
        return None
    except Exception as e:
        logging.error("Unexpected error fetching image: %s", e)
        return None

def convert_yn_to_yes_no(value):
    if isinstance(value, str):
        if value.lower() == 'y':
            return 'Yes'
        elif value.lower() == 'n':
            return 'No'
    return value

def parse_and_format_date(date_str):
    """
    For regular dates (dob, exit_date, orientation_date, last_attended)
    -> format as m/d/yyyy (no leading zeros)
    """
    s = _normalize_date_string(date_str)
    if not s:
        return ''
    # try a couple explicit formats first, then fall back to pandas
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m/%d/%y'):
        try:
            d = datetime.strptime(s, fmt)
            return f"{d.month}/{d.day}/{d.year}"
        except ValueError:
            pass
    try:
        d = pd.to_datetime(s, errors='raise', dayfirst=False)
        return f"{d.month}/{d.day}/{d.year}"
    except Exception:
        return ''


def parse_and_format_p_a_date(date_str):
    """
    For P1..P27 and A1..A18 -> format as mm/dd/yy (with leading zeros)
    """
    s = _normalize_date_string(date_str)
    if not s:
        return ''
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m/%d/%y'):
        try:
            d = datetime.strptime(s, fmt)
            return d.strftime('%m/%d/%y')
        except ValueError:
            pass
    try:
        d = pd.to_datetime(s, errors='raise', dayfirst=False)
        return d.strftime('%m/%d/%y')
    except Exception:
        return ''

import re

def _normalize_date_string(s: str) -> str:
    if not isinstance(s, str):
        return ''
    # normalize dashes, strip ALL internal whitespace/newlines, trim commas
    s = s.replace('\u2013', '-').replace('\u2014', '-')
    s = re.sub(r'\s+', '', s)
    return s.strip().strip(',')


# Define function to determine fee problems
def determine_fee_problems(balance):
    try:
        balance = float(balance)
        return "No" if balance == 0 else "Yes"
    except ValueError:
        logging.error(f"Invalid balance value: {balance}")
        return "Unknown"

# Define function to format balance
def format_balance(balance):
    try:
        balance = float(balance)
        return f"${int(balance)}"  # Convert to whole number and format as currency
    except ValueError:
        return "$0"  # Default to $0 if invalid

def determine_stage_of_change(attended_str, required_str):
    try:
        attended = float(attended_str)
        required = float(required_str)
    except ValueError:
        logging.error(f"Invalid attended/required values: attended={attended_str}, required={required_str}")
        return 'Alert'

    if required == 0:
        return 'Alert'

    # Exact logic for 18-week program
    if required == 18:
        if 1 <= attended <= 4:
            return 'Precontemplation'
        elif 5 <= attended <= 8:
            return 'Contemplation'
        elif 9 <= attended <= 11:
            return 'Preparation'
        elif 12 <= attended <= 15:
            return 'Action'
        elif 16 <= attended <= 18:
            return 'Maintenance'
        else:
            logging.warning(f"Unexpected attendance value for BIPP 18-week: attended={attended}")
            return 'Alert'

    # Exact logic for 27-week program
    if required == 27:
        if 1 <= attended <= 7:
            return 'Precontemplation'
        elif 8 <= attended <= 14:
            return 'Contemplation'
        elif 15 <= attended <= 18:
            return 'Preparation'
        elif 19 <= attended <= 23:
            return 'Action'
        elif 24 <= attended <= 27:
            return 'Maintenance'
        else:
            logging.warning(f"Unexpected attendance value for BIPP 27-week: attended={attended}")
            return 'Alert'

    # Adaptive logic for other durations using 18-week proportions
    ratio = attended / required
    # Using the 18-week thresholds proportionally:
    # Precontemplation: ≤ 4/18 ≈ 0.2222
    # Contemplation: ≤ 8/18 ≈ 0.4444
    # Preparation: ≤ 11/18 ≈ 0.6111
    # Action: ≤ 15/18 ≈ 0.8333
    # Maintenance: > 0.8333

    if ratio > 0 and ratio <= (4/18):
        return 'Precontemplation'
    elif ratio <= (8/18):
        return 'Contemplation'
    elif ratio <= (11/18):
        return 'Preparation'
    elif ratio <= (15/18):
        return 'Action'
    elif ratio <= 1.0:
        return 'Maintenance'
    else:
        logging.warning(f"Ratio out of expected range for required={required}, attended={attended}")
        return 'Alert'

def fill_placeholders(row):
    # Set gender placeholders
    gender = row.get('gender', '').lower()
    if gender == 'male':
        row['gender1'] = 'his'
        row['gender2'] = 'he'
        row['gender3'] = "Men's"
        row['gender4'] = 'him'
        row['gender5'] = "He"
        row['gender6'] = 'himself'
        row['gender7'] = "Mr."
    elif gender == 'female':
        row['gender1'] = 'her'
        row['gender2'] = 'she'
        row['gender3'] = "Women's"
        row['gender4'] = 'her'
        row['gender5'] = "She"
        row['gender6'] = 'herself'
        row['gender7'] = "Ms."
    else:
        for k,v in {
            'gender1':'their','gender2':'they','gender3':"Participants",
            'gender4':'them','gender5':'They','gender6':'themself','gender7':''
        }.items():
            row[k] = v


    # Parse and format key dates
    for date_field in ['dob', 'exit_date', 'orientation_date', 'last_attended']:
        if date_field in row:
            row[date_field] = parse_and_format_date(row[date_field])

    # Inside fill_placeholders or wherever you handle attendance placeholders:
    # Clear any old values first:
    for i in range(1, 31):
        row[f'D{i}'] = ''

    # Fill placeholders up to the correct count:
    attended_str = row.get('attended', '0')
    required_str = row.get('required_sessions', '0')

    try:
        attended = float(attended_str)
        required = float(required_str)
    except ValueError:
        attended = 0
        required = 0

    if required > 0:
        # Determine how many of the D1..D30 placeholders should get the checkmark
        placeholders_to_fill = int((attended / required) * 30)
        for i in range(1, placeholders_to_fill + 1):
            row[f'D{i}'] = '\u2714'   # Heavy checkmark: '✔'


    # Replace '&' with 'and'
    if 'client_note' in row:
        row['client_note'] = row['client_note'].replace('&', 'and')

    # Format P1-P27 and A1-A18 as mm/dd/yy (format both cases so template gets the right one)
    for i in range(1, 28):
        for key in (f'p{i}', f'P{i}'):
            if key in row:
                row[key] = parse_and_format_p_a_date(row[key])

    for i in range(1, 19):
        for key in (f'a{i}', f'A{i}'):
            if key in row:
                row[key] = parse_and_format_p_a_date(row[key])


    # Determine client stage of change
    row['client_stagechange'] = determine_stage_of_change(attended, required)

    return row

def sanitize_filename(filename):
    invalid_chars = '<>:"/\\|?*'
    for char in invalid_chars:
        filename = filename.replace(char, '')
    return filename.strip()

def ensure_columns(df, cols, default=''):
    for c in cols:
        if c not in df.columns:
            df[c] = default
    return df

def _is_yes(v):
    return str(v).strip().lower() in ('yes','y','1','true','✓','check','checked')

# AFTER  (✓ / ✖)
def _chk(ok):
    return '\u2714' if ok else '\u2716'


def _dwag_exit_derived(row):
    # attendance
    absences = 0
    try: absences = int(float(str(row.get('absence_unexcused','0') or 0)))
    except: pass
    required = 0
    try: required = float(str(row.get('required_sessions','0') or 0))
    except: pass

    has_orientation = bool(str(row.get('orientation_date','')).strip())
    row['attended_initial_assessment']   = _chk(has_orientation)
    row['attended_initial_assessessment']= row['attended_initial_assessment']  # template typo

    if required:
        ratio = (required - absences) / required
        row['consistent_attendance'] = _chk(ratio >= 0.90)   # within 10% unexcused
    else:
        row['consistent_attendance'] = _chk(False)

    # text, not a checkbox
    row['three_or_more_absences'] = (
        f"{absences} absences" if absences >= 3 else ("No" if absences == 0 else "Within 3")
    )

    # conduct / agreement
    respectful   = _is_yes(row.get('respectful_to_group','No'))
    disruptive   = _is_yes(row.get('disruptive_argumentitive', row.get('disruptive_arumentitive','No')))
    humor_bad    = _is_yes(row.get('humor_inappropriate','No'))
    blames       = _is_yes(row.get('blames_victim','No'))
    appears      = _is_yes(row.get('appears_drug_alcohol','No'))
    staff_bad    = _is_yes(row.get('inappropriate_to_staff','No'))
    dialogue     = _is_yes(row.get('speaks_significantly_in_group','No'))
    takes_resp   = _is_yes(row.get('takes_responsibility_for_past','No'))

    row['followed_group_rules']             = _chk(respectful and not disruptive and not humor_bad and not staff_bad)
    row['adhered_to_participant_agreement'] = _chk(not appears and not blames and not staff_bad)
    row['participated_in_dialogue']         = _chk(dialogue or (respectful and not disruptive))
    row['participated_in_dialoge']          = row['participated_in_dialogue']  # template typo
    row['accountability_consistent']        = _chk(takes_resp)
    return row

# Load CSV
try:
    df = pd.read_csv(csv_file_path)
    logging.info(f"Loaded {len(df)} records from CSV.")
except Exception as e:
    logging.error(f"Failed to load CSV file: {e}")
    sys.exit(1)

# Don’t assign '' into numeric dtypes; replace NaNs first, then cast
df = df.where(pd.notna(df), '')     # no dtype clash
df = df.astype(str)                 # now everything is string for templating

# make sure columns we read are present (safe defaults)
df = ensure_columns(df, ['exit_date','balance','last_attended','orientation_date'], '')
df = ensure_columns(df, ['absence_unexcused','required_sessions','attended'], '0')
df = ensure_columns(df, [
    'orientation_date','respectful_to_group','disruptive_argumentitive','disruptive_arumentitive',
    'humor_inappropriate','blames_victim','appears_drug_alcohol','inappropriate_to_staff',
    'speaks_significantly_in_group','takes_responsibility_for_past'
], '')

# Convert exit_date to datetime, filter, then convert back to string
df['exit_date'] = pd.to_datetime(df['exit_date'], errors='coerce')
df = df.dropna(subset=['exit_date'])

df = df[df['program_name'].isin(['BIPP (male)', 'BIPP (female)'])]
df = df[df['exit_reason'].isin(['Violation of Requirements', 'Unable to Participate', 'Death', 'Moved'])]
df = df.sort_values(by='exit_date', ascending=False)

if start_date_str and end_date_str:
    try:
        start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
        end_date = datetime.strptime(end_date_str, '%Y-%m-%d')
        df = df[(df['exit_date'] >= start_date) & (df['exit_date'] <= end_date)]
    except Exception as e:
        logging.error(f"Error during date filtering: {e}")
        sys.exit(1)

if df.empty:
    logging.warning("No records found after filtering. No documents will be generated.")
    sys.exit(0)

# Convert exit_date back to string for parsing
df['exit_date'] = df['exit_date'].dt.strftime('%Y-%m-%d')

# Handle last_attended
def fix_last_attended(row):
    if not row['last_attended'].strip():
        return row['orientation_date']
    return row['last_attended']

df['last_attended'] = df.apply(fix_last_attended, axis=1)

# Apply determine_fee_problems and format_balance
df['feeproblems'] = df['balance'].apply(determine_fee_problems)
df['balance'] = df['balance'].apply(format_balance)

# -------------------------------------------------------------------
# Derive extra compliance / benchmark fields for DWAG template
# -------------------------------------------------------------------
# ---------------------------------------------------------------------------
# 1)  NORMALISE BOOLEAN / ENUM COLUMNS  (run once, right after df.astype(str))
# ---------------------------------------------------------------------------
# camelCase → snake_case expected by the template helpers
df = df.rename(columns={
    'speaksSignificantlyInGroup'        : 'speaks_significantly_in_group',
    'respectfulTowardsGroup'            : 'respectful_to_group',
    'takesResponsibilityForPastBehavior': 'takes_responsibility_for_past',
    'disruptiveOrArgumentitive'         : 'disruptive_argumentitive',
    'inappropriateHumor'                : 'humor_inappropriate',
    'blamesVictim'                      : 'blames_victim',
    'drug_alcohol'                      : 'appears_drug_alcohol',
    'inappropriate_behavior_to_staff'   : 'inappropriate_to_staff'
})

# convert 1/0 (or y/n) to the “Yes”/“No” strings used everywhere else
BOOL_COLS = [
    'speaks_significantly_in_group', 'respectful_to_group',
    'takes_responsibility_for_past', 'disruptive_argumentitive',
    'humor_inappropriate', 'blames_victim',
    'appears_drug_alcohol', 'inappropriate_to_staff'
]
for col in BOOL_COLS:
    if col not in df.columns:
        df[col] = 'No'
    else:
        df[col] = df[col].map({'1':'Yes','0':'No','y':'Yes','n':'No','Y':'Yes','N':'No'}).fillna('No')


# ---------------------------------------------------------------------------
# 2)  BENCHMARKS & FLAGS FOR PROGRESS-REPORT + EXIT NOTICE
# ---------------------------------------------------------------------------
ABSENCE_THRESHOLD = 0.10       # < 10 % un-excused absences = “consistent”
ABSENCE_LIMIT     = 3          # termination if ≥ 3 un-excused

def derive_benchmark_flags(row):
    # ------- attendance -------
    row['attended_initial_assessment'] = 'Yes' if row['orientation_date'] else 'No'

    try:
        attended   = float(row.get('attended', 0))
        unexcused  = float(row.get('absence_unexcused', 0))
        required   = float(row.get('required_sessions', 0))
    except ValueError:
        attended = unexcused = required = 0

    if required:
        ratio = (required - unexcused) / required
        row['consistent_attendance'] = 'Yes' if ratio >= (1 - ABSENCE_THRESHOLD) else 'No'
    else:
        row['consistent_attendance'] = 'No'

    row['three_or_more_absences'] = 'Yes' if unexcused >= ABSENCE_LIMIT else 'No'

    # ------- conduct / rule-compliance -------
    good_flags = all([
        row['speaks_significantly_in_group'] == 'Yes',
        row['respectful_to_group']           == 'Yes',
        row['takes_responsibility_for_past'] == 'Yes',
        row['disruptive_argumentitive']      == 'No',
        row['humor_inappropriate']           == 'No',
        row['blames_victim']                 == 'No',
    ])
    row['followed_group_rules'] = 'Yes' if good_flags else 'No'

    agreement_flags = all([
        row['appears_drug_alcohol'] == 'No',
        row['inappropriate_to_staff'] == 'No',
    ])
    row['adhered_to_participant_agreement'] = 'Yes' if agreement_flags else 'No'

    # ------- shortcuts reused in template -------
    row['participated_in_dialogue'] = row['speaks_significantly_in_group']
    row['accountability_consistent'] = row['takes_responsibility_for_past']
    return row

if clinic_folder.lower() == 'dwag':
    # Use ✓ / ✖ fields for DWAG
    df = df.apply(_dwag_exit_derived, axis=1)
else:
    # Other clinics keep Yes/No flags
    df = df.apply(derive_benchmark_flags, axis=1)


# ---------------------------------------------------------------------------
# 3)  BUILD ONE MERGED “EXIT REASONS” STRING
# ---------------------------------------------------------------------------
def build_exit_reasons(row):
    reasons = []

    # database-supplied exit_reason
    if row.get('exit_reason'):
        reasons.append(row['exit_reason'])

    # ≥ 3 un-excused absences
    try:
        unexcused = int(float(row.get('absence_unexcused', 0)))
    except ValueError:
        unexcused = 0
    if unexcused >= ABSENCE_LIMIT:
        reasons.append(f"{unexcused} un-excused absences")

    # conduct problems
    bad_map = {
        'disruptive_argumentitive'      : "Disruptive / argumentative in group",
        'humor_inappropriate'           : "Inappropriate humor",
        'blames_victim'                 : "Blamed victim for abuse",
        'respectful_to_group'           : "Disrespectful toward group members",
        'speaks_significantly_in_group' : "Did not participate in dialogue",
        'takes_responsibility_for_past' : "Refused accountability for past abuse",
        'appears_drug_alcohol'          : "Appeared under influence in group",
        'inappropriate_to_staff'        : "Inappropriate behavior toward staff",
    }
    for col, msg in bad_map.items():
        val = row[col]
        trigger = (val == 'Yes') if col not in ['respectful_to_group',
                                                'speaks_significantly_in_group',
                                                'takes_responsibility_for_past'] else (val == 'No')
        if trigger:
            reasons.append(msg)

    row['exit_reasons_block'] = "; ".join(reasons) if reasons else "See case notes"
    return row

df = df.apply(build_exit_reasons, axis=1)


# ---------------------------------------------------------------------------
# 4)  REMAINING PLACEHOLDERS
# ---------------------------------------------------------------------------
df = df.apply(fill_placeholders, axis=1)



# Convert Y/N to Yes/No
columns_to_convert = [
    "speaks_significantly_in_group", "respectful_to_group",
    "takes_responsibility_for_past", "disruptive_argumentitive",
    "humor_inappropriate", "blames_victim",
    "appears_drug_alcohol", "inappropriate_to_staff"
]

if clinic_folder.lower() != 'dwag':
    columns_to_convert = [
        "speaks_significantly_in_group","respectful_to_group","takes_responsibility_for_past",
        "disruptive_argumentitive","humor_inappropriate","blames_victim",
        "appears_drug_alcohol","inappropriate_to_staff"
    ]


for col in columns_to_convert:
    if col in df.columns:
        df[col] = df[col].apply(convert_yn_to_yes_no)


def render_and_save(template_path, doc_type, context, output_dir):
    try:
        doc = DocxTemplate(template_path)

        # Insert image if available
        if 'image_url' in context and context['image_url']:
            image = fetch_image(context['image_url'], doc=doc)
            context['image'] = image if image else ''

        doc.render(context)
        align_tables_left(doc)

        first_name = sanitize_filename(context.get('first_name', ''))
        last_name = sanitize_filename(context.get('last_name', ''))
        docx_filename = f"{last_name} {first_name} {doc_type}.docx"
        docx_path = os.path.join(generated_documents_dir, docx_filename)

        doc.save(docx_path)
        time.sleep(1)

        pdf_path = docx_to_pdf(docx_path, generated_documents_dir)
        if pdf_path:
            os.remove(docx_path)
            return pdf_path
        else:
            logging.error(f"PDF conversion failed for {docx_path}")
            return None
    except Exception as e:
        logging.error(f"Error rendering {doc_type}: {e}")
        record_error(f"Render failed for {doc_type} / {context.get('last_name','')}, {context.get('first_name','')}: {e}")
        return None


def generate_documents(row, output_dir_path):
    context = row.to_dict()

    # Add report_date (today's date) in m/d/yyyy format no leading zeros
    now = datetime.now()
    context['report_date'] = f"{now.month}/{now.day}/{now.year}"

    # <<< INSERT THIS BLOCK HERE >>>
    report_template = progress_report_template_path
    report_suffix   = "Exit PReport"
    group_name = str(context.get("group_name", "")).lower()
    if ("virtual" in group_name) or ("women" in group_name):
        report_template = virtual_progress_report_template_path
        report_suffix   = "Exit PReport Vir"
    # <<< END INSERT >>>

    # Generate Exit Notice
    exit_pdf = render_and_save(exit_notice_template_path, "BIPP Exit Notice", context, output_dir_path)
    # Generate Exit Progress Report
    progress_pdf = render_and_save(report_template, f"BIPP {report_suffix}", context, output_dir_path)

    # Organize PDFs by case manager
    # Organize PDFs by case manager / facilitator / instructor (clinic differences)
    def pick(d, *keys, default=''):
        for k in keys:
            v = d.get(k, '')
            if isinstance(v, str) and v.strip():
                return v.strip()
        return default

    try:
        raw_office = pick(row, 'case_manager_office','facilitator_office','instructor_office', default='')
        office_dir = (output_dir_path if not raw_office or raw_office.lower() in ('nan','none','null','')
                    else os.path.join(output_dir_path, sanitize_filename(raw_office)))

        raw_first = pick(row, 'case_manager_first_name','facilitator_first_name','instructor_first_name', default='')
        raw_last  = pick(row, 'case_manager_last_name', 'facilitator_last_name', 'instructor_last_name',  default='')
        manager_subfolder = (f"{sanitize_filename(raw_first)} {sanitize_filename(raw_last)}".strip()
                            or (sanitize_filename(raw_last) if raw_last else 'Unknown Manager'))

        manager_dir = os.path.join(office_dir, manager_subfolder)
        os.makedirs(manager_dir, exist_ok=True)

        for pdf_file_path in [exit_pdf, progress_pdf]:
            if pdf_file_path and os.path.exists(pdf_file_path):
                shutil.move(pdf_file_path, os.path.join(manager_dir, os.path.basename(pdf_file_path)))
                logging.info(f"Moved {pdf_file_path} to {manager_dir}")
            else:
                logging.warning("PDF file not found or not generated.")
    except Exception as e:
        logging.error(f"Error organizing generated PDFs: {e}")


processed = 0
skipped   = 0

try:
    for idx, row in df.iterrows():
        try:
            generate_documents(row, central_location_path)
            processed += 1
        except Exception as row_e:
            skipped += 1
            record_error(f"[row {idx}] Unexpected error: {row_e}")
            continue
except Exception as e:
    record_error(f"[FATAL] Loop crashed: {e}")

logging.info(f"BIPP Exit Notices completed. Processed={processed}, Skipped={skipped}")
print(f"BIPP Exit Notices complete. Processed={processed}, Skipped={skipped}")
sys.exit(0)
