#!/usr/bin/env python3
# check_balances.py — totals tables + per-client blocks + ledger details
# - Uses `ledger_date` from ledger CSV
# - Darker shading for DEBIT rows
# - One client per page
# - Compact spacing between summary and ledger tables
#
# Usage:
#   python3 check_balances.py --csv_file <path>
#                             --output_directory <dir>
#                             [--report_date "MM/DD/YYYY"]
#                             [--ledger_csv <path>]   # CSV: client_id, ledger_date, amount, note

import argparse
import os
import subprocess
from datetime import datetime

import pandas as pd
from docx import Document
from docx.shared import Inches
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ISSUE_ORDER = [
    "Owes for attended",
    "No payments yet",
    "Missing/zero fee",
    "Missing required_sessions",
    "Overpaid",
]


def parse_args():
    ap = argparse.ArgumentParser(description="Generate balance report PDF from issue CSV.")
    ap.add_argument("--csv_file", required=True, help="Input CSV produced by check_balances.php")
    ap.add_argument("--ledger_csv", required=False, default="", help="Optional ledger lines CSV with dates and notes")
    ap.add_argument("--report_date", required=False, default=datetime.now().strftime("%m/%d/%Y"))
    ap.add_argument("--output_directory", required=True, help="Directory for output .docx/.pdf")
    return ap.parse_args()


def load_ledger_df(path: str) -> pd.DataFrame:
    if not path or not os.path.exists(path):
        return pd.DataFrame(columns=["client_id", "ledger_date", "amount", "note"])
    df = pd.read_csv(path)

    # Normalize columns
    for col, default in [
        ("client_id", None),
        ("ledger_date", ""),
        ("amount", 0.0),
        ("note", ""),
    ]:
        if col not in df.columns:
            df[col] = default

    df["client_id"] = pd.to_numeric(df["client_id"], errors="coerce").astype("Int64")
    df["amount"] = pd.to_numeric(df["amount"], errors="coerce").fillna(0.0)

    # Format ledger_date as MM/DD/YYYY when possible
    df["ledger_date"] = pd.to_datetime(df["ledger_date"], errors="coerce").dt.strftime("%m/%d/%Y")
    df["ledger_date"] = df["ledger_date"].fillna("")

    # Type label for shading
    df["type"] = df["amount"].apply(lambda a: "Credit" if a > 0 else ("Debit" if a < 0 else "Zero"))
    return df


def money(x):
    try:
        return f"${float(x):,.2f}"
    except Exception:
        return str(x)


def num(x):
    try:
        v = float(x)
        return str(int(v)) if v.is_integer() else f"{v:.2f}"
    except Exception:
        return str(x)


def normalize_df(df: pd.DataFrame) -> pd.DataFrame:
    expected = [
        "client_id",
        "first_name",
        "last_name",
        "program_id",
        "program_name",
        "fee",
        "required_sessions",
        "attended_sessions",
        "total_paid",
        "expected_paid_to_date",
        "current_balance",
        "total_expected_by_graduation",
        "issue",
    ]
    missing = [c for c in expected if c not in df.columns]
    if missing:
        raise ValueError(f"Missing required columns in CSV: {', '.join(missing)}")

    for c in [
        "fee",
        "required_sessions",
        "attended_sessions",
        "total_paid",
        "expected_paid_to_date",
        "current_balance",
        "total_expected_by_graduation",
    ]:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)

    for c in ["first_name", "last_name", "program_name", "issue"]:
        df[c] = df[c].fillna("").astype(str)

    df["program_label"] = df.apply(
        lambda r: r["program_name"].strip() if r["program_name"].strip() else f"Program {int(r['program_id'])}",
        axis=1,
    )
    df["name_disp"] = df.apply(lambda r: f"{r['last_name'].strip()}, {r['first_name'].strip()}".strip(", "), axis=1)
    df["__last"] = df["last_name"].str.strip().str.lower()
    df["client_id"] = pd.to_numeric(df["client_id"], errors="coerce").astype("Int64")
    return df


def explode_issues(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, r in df.iterrows():
        issues = [s.strip() for s in str(r["issue"]).split("|") if s.strip()]
        if not issues:
            continue
        for it in issues:
            rr = r.to_dict()
            rr["issue_single"] = it
            rows.append(rr)
    cols = list(df.columns) + ["issue_single"]
    return pd.DataFrame(rows, columns=cols) if rows else pd.DataFrame(columns=cols)


def compute_program_totals(df: pd.DataFrame) -> pd.DataFrame:
    base = df.groupby("program_label", dropna=False)
    agg = base.agg(
        clients=("client_id", "nunique"),
        attended=("attended_sessions", "sum"),
        expected_to_date=("expected_paid_to_date", "sum"),
        paid=("total_paid", "sum"),
        total_expected=("total_expected_by_graduation", "sum"),
        bal_pos=("current_balance", lambda s: float(s[s > 0].sum())),
        bal_neg=("current_balance", lambda s: float(s[s < 0].sum())),  # negative
    ).reset_index()

    under = df[df["current_balance"] > 0].groupby("program_label")["client_id"].nunique()
    over = df[df["current_balance"] < 0].groupby("program_label")["client_id"].nunique()
    agg["under_cnt"] = agg["program_label"].map(under).fillna(0).astype(int)
    agg["over_cnt"] = agg["program_label"].map(over).fillna(0).astype(int)

    agg["bal_neg_abs"] = agg["bal_neg"].abs()
    agg["net_bal"] = agg["bal_pos"] - agg["bal_neg_abs"]
    return agg.sort_values("program_label", kind="stable")


# ----- DOCX helpers -----
def shade_cell(cell, fill_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def shade_row(cells, fill_hex: str):
    for c in cells:
        shade_cell(c, fill_hex)


def align_right(cell):
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def set_tbl_fixed_layout(table):
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def set_col_widths(table, widths_in_inches):
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in_inches):
            cell.width = Inches(w)


# ----- Rendering blocks -----
def add_totals_tables(doc, totals_df):
    # A) Participation & money to date
    doc.add_heading("Totals by Program — Participation & Money to Date", level=2)
    headers_a = ["Program", "Clients", "Attended", "ExpTD", "Paid"]
    t_a = doc.add_table(rows=1, cols=len(headers_a))
    t_a.style = "Table Grid"
    set_tbl_fixed_layout(t_a)
    set_col_widths(t_a, [2.30, 0.85, 0.95, 1.20, 1.20])

    for i, h in enumerate(headers_a):
        c = t_a.cell(0, i)
        c.text = h
        if i > 0:
            align_right(c)

    for _, r in totals_df.iterrows():
        rw = t_a.add_row().cells
        rw[0].text = str(r["program_label"])
        rw[1].text = num(r["clients"]);              align_right(rw[1])
        rw[2].text = num(r["attended"]);             align_right(rw[2])
        rw[3].text = money(r["expected_to_date"]);   align_right(rw[3])
        rw[4].text = money(r["paid"]);               align_right(rw[4])

    doc.add_paragraph()  # small spacer

    # B) Balance snapshot
    doc.add_heading("Totals by Program — Balance Snapshot", level=2)
    headers_b = ["Program", "Under#", "Over#", "Bal+", "Bal−", "Net", "GradTot"]
    t_b = doc.add_table(rows=1, cols=len(headers_b))
    t_b.style = "Table Grid"
    set_tbl_fixed_layout(t_b)
    set_col_widths(t_b, [2.30, 0.80, 0.80, 1.05, 1.05, 1.05, 1.20])

    for i, h in enumerate(headers_b):
        c = t_b.cell(0, i)
        c.text = h
        if i > 0:
            align_right(c)

    for _, r in totals_df.iterrows():
        rw = t_b.add_row().cells
        rw[0].text = str(r["program_label"])
        rw[1].text = num(r["under_cnt"]);            align_right(rw[1])
        rw[2].text = num(r["over_cnt"]);             align_right(rw[2])
        rw[3].text = money(r["bal_pos"]);            align_right(rw[3])
        rw[4].text = money(r["bal_neg_abs"]);        align_right(rw[4])
        rw[5].text = money(r["net_bal"]);            align_right(rw[5])
        rw[6].text = money(r["total_expected"]);     align_right(rw[6])

    doc.add_paragraph(
        "Legend: Under#=clients owing; Over#=clients overpaid; ExpTD=expected to date; "
        "Bal+=sum owed; Bal−=sum overpaid; Net=Bal+−Bal−; GradTot=total expected by graduation."
    )


def add_client_summary_table(doc: Document, r: pd.Series):
    # Summary line
    name = r["name_disp"]
    bal = float(r["current_balance"])
    if bal > 0:
        summary = f"{name} owes {money(bal)}."
    elif bal < 0:
        summary = f"{name} is overpaid by {money(abs(bal))}."
    else:
        summary = f"{name} is paid up to date."
    p = doc.add_paragraph()
    p.add_run(summary).bold = True

    # Mini table: Fee | Attended | ExpTD | Paid | Balance | GradTot
    headers = ["Fee", "Attended", "ExpTD", "Paid", "Balance", "GradTot"]
    values = [
        money(r["fee"]),
        num(r["attended_sessions"]),
        money(r["expected_paid_to_date"]),
        money(r["total_paid"]),
        money(r["current_balance"]),
        money(r["total_expected_by_graduation"]),
    ]
    tbl = doc.add_table(rows=2, cols=6)
    tbl.style = "Light Shading Accent 1"
    set_tbl_fixed_layout(tbl)
    set_col_widths(tbl, [0.9, 0.9, 1.0, 1.0, 1.1, 1.1])

    for i, h in enumerate(headers):
        hdr = tbl.cell(0, i)
        hdr.text = h
        align_right(hdr)

    for i, v in enumerate(values):
        c = tbl.cell(1, i)
        c.text = v
        align_right(c)

    # Balance shading
    bal_cell = tbl.cell(1, 4)
    if bal > 0:
        shade_cell(bal_cell, "F8D7DA")  # light red
    elif bal < 0:
        shade_cell(bal_cell, "D4EDDA")  # light green

    # very small spacer
    doc.add_paragraph().add_run("")


def add_client_ledger_table(doc: Document, ledger_rows: pd.DataFrame):
    """
    Ledger detail table under each client:
      Columns: Date | Type | Amount | Note
      Type ∈ {Credit, Debit, Zero}. Debits shaded darker.
    """
    if ledger_rows is None or ledger_rows.empty:
        return

    tmp = ledger_rows.copy()
    tmp["amount"] = pd.to_numeric(tmp["amount"], errors="coerce").fillna(0.0)
    if "ledger_date" not in tmp.columns:
        tmp["ledger_date"] = ""
    if "note" not in tmp.columns:
        tmp["note"] = ""
    if "type" not in tmp.columns:
        tmp["type"] = tmp["amount"].apply(lambda a: "Credit" if a > 0 else ("Debit" if a < 0 else "Zero"))

    headers = ["Date", "Type", "Amount", "Note"]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    set_tbl_fixed_layout(t)
    set_col_widths(t, [1.1, 0.9, 1.0, 3.8])

    for i, h in enumerate(headers):
        t.cell(0, i).text = h
        if h == "Amount":
            align_right(t.cell(0, i))

    for _, rr in tmp.iterrows():
        amt = float(rr["amount"])
        typ = rr.get("type", "")
        rw = t.add_row().cells
        rw[0].text = str(rr.get("ledger_date", "")).strip()
        rw[1].text = typ
        rw[2].text = money(amt); align_right(rw[2])
        rw[3].text = str(rr.get("note", "")).strip()

        # Shading
        if typ == "Debit":
            shade_row(rw, "E99A9A")   # darker red
        elif typ == "Zero":
            shade_row(rw, "EEEEEE")

    doc.add_paragraph().add_run("")


def render_details_by_issue(doc: Document,
                            df_issues: pd.DataFrame,
                            df_base: pd.DataFrame,
                            ledger_df: pd.DataFrame | None):
    doc.add_heading("Details by Issue", level=2)
    first_cat = False

    # index ledger by client
    ledger_by_client = {}
    if ledger_df is not None and not ledger_df.empty and "client_id" in ledger_df.columns:
        ledger_df["client_id"] = pd.to_numeric(ledger_df["client_id"], errors="coerce").astype("Int64")
        for cid, g in ledger_df.groupby("client_id"):
            ledger_by_client[int(cid)] = g.copy()

    for issue in ISSUE_ORDER:
        grp_issue = df_issues[df_issues["issue_single"] == issue]
        if grp_issue.empty:
            continue

        # Page break between categories
        if first_cat:
            doc.add_page_break()
        first_cat = True

        doc.add_heading(issue, level=3)

        for program_label, grp in grp_issue.groupby("program_label"):
            doc.add_heading(str(program_label), level=4)
            grp = grp.sort_values(["__last", "first_name"], kind="stable")

            for _, r in grp.iterrows():
                add_client_summary_table(doc, r)
                cid = int(r["client_id"]) if pd.notna(r["client_id"]) else None
                ldf = ledger_by_client.get(cid)
                add_client_ledger_table(doc, ldf)
                doc.add_page_break()  # one client per page


def save_results_to_doc(df_base: pd.DataFrame,
                        df_issues: pd.DataFrame,
                        ledger_df: pd.DataFrame | None,
                        report_date: str,
                        output_path_docx: str):
    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.page_height = Inches(11)
    section.page_width = Inches(8.5)
    section.orientation = WD_ORIENT.PORTRAIT

    # Header
    doc.add_heading("Client Balance Report", level=1)
    doc.add_paragraph(f"Report date: {report_date}")
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    if df_base.empty:
        doc.add_paragraph("No clients with payment issues found.")
        doc.save(output_path_docx)
        return

    # Program totals
    totals = compute_program_totals(df_base)
    add_totals_tables(doc, totals)

    # Grand totals
    g_clients = int(df_base["client_id"].nunique())
    g_att = int(df_base["attended_sessions"].sum())
    g_exp_to_date = float(df_base["expected_paid_to_date"].sum())
    g_paid = float(df_base["total_paid"].sum())
    g_pos = float(df_base.loc[df_base["current_balance"] > 0, "current_balance"].sum())
    g_neg_abs = float(-df_base.loc[df_base["current_balance"] < 0, "current_balance"].sum())
    g_net = g_pos - g_neg_abs
    g_total_expected = float(df_base["total_expected_by_graduation"].sum())

    doc.add_paragraph(
        f"Grand totals — Clients: {g_clients} | Attended: {g_att} | "
        f"ExpTD: {money(g_exp_to_date)} | Paid: {money(g_paid)} | "
        f"Bal+: {money(g_pos)} | Bal−: {money(g_neg_abs)} | Net: {money(g_net)} | "
        f"GradTot: {money(g_total_expected)}"
    )

    # Details with per-client page breaks
    doc.add_page_break()
    render_details_by_issue(doc, df_issues, df_base, ledger_df)

    doc.save(output_path_docx)


def convert_docx_to_pdf(docx_path: str, pdf_path: str):
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            os.path.dirname(pdf_path),
            docx_path,
        ],
        check=True,
    )


def process_balances(csv_file: str,
                     ledger_csv: str,
                     report_date: str,
                     output_directory: str) -> str:
    df = pd.read_csv(csv_file)
    df = normalize_df(df)
    df_i = explode_issues(df)

    ledger_df = load_ledger_df(ledger_csv) if ledger_csv else pd.DataFrame()

    today = datetime.now().strftime("%Y%m%d")
    out_docx = os.path.join(output_directory, f"BalanceReport_{today}.docx")
    out_pdf = os.path.join(output_directory, f"BalanceReport_{today}.pdf")
    os.makedirs(output_directory, exist_ok=True)

    save_results_to_doc(df, df_i, ledger_df, report_date, out_docx)
    convert_docx_to_pdf(out_docx, out_pdf)
    try:
        os.remove(out_docx)
    except Exception:
        pass
    return out_pdf


def main():
    args = parse_args()
    os.makedirs(args.output_directory, exist_ok=True)
    try:
        pdf_path = process_balances(args.csv_file, args.ledger_csv, args.report_date, args.output_directory)
        print(f"Balance report generated: {pdf_path}")
    except Exception as e:
        print(f"Error generating balance report: {e}")


if __name__ == "__main__":
    main()
